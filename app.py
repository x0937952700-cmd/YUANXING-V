
from flask import Flask, render_template, request, jsonify, session, redirect, url_for, Response, stream_with_context, send_file, send_from_directory
from datetime import timedelta, datetime
from functools import wraps
import os
import io
import time
import hashlib
import json
import re
from PIL import Image
from werkzeug.utils import secure_filename
from openpyxl import Workbook

from db import (
    init_db, get_user, create_user, update_password, log_action,
    save_inventory_item, list_inventory, save_order, save_master_order,
    ship_order, preview_ship_order, get_shipping_records, save_correction, log_error,
    save_image_hash, image_hash_exists, upsert_customer, get_customers,
    get_customer, warehouse_get_cells, warehouse_save_cell, warehouse_move_item, warehouse_add_column,
    warehouse_add_slot, warehouse_remove_slot,
    inventory_summary, warehouse_summary, list_backups, get_orders, get_master_orders,
    list_users, set_user_blocked, get_setting, set_setting, verify_password, row_to_dict, get_db, sql, rows_to_dict, fetchone_dict, now,
    register_submit_request, list_corrections_rows, delete_correction, save_customer_alias, list_customer_aliases, delete_customer_alias,
    record_recent_slot, get_recent_slots, add_audit_trail, list_audit_trails, get_customer_spec_stats, update_customer_item, update_items_material, delete_customer_item,
    create_todo_item, list_todo_items, get_todo_item, delete_todo_item, complete_todo_item, restore_todo_item, reorder_todo_items,
    delete_customer, get_customer_relation_counts, get_customer_by_uid, restore_customer, effective_product_qty, product_display_size, product_support_text, product_sort_tuple, format_product_text_height2, clean_material_value
)
from ocr import parse_ocr_text, process_native_ocr_text, clean_ocr_noise
from backup import run_daily_backup

app = Flask(__name__)
# FIX52：優先使用 Render 環境變數 SECRET_KEY。
# 若尚未設定，改用 DATABASE_URL 雜湊產生穩定 fallback，避免每次重啟都登出。
_SECRET_KEY = os.getenv("SECRET_KEY") or ("stable-" + hashlib.sha256((os.getenv("DATABASE_URL", "yuanxing-local") + "|yuanxing-fix53").encode("utf-8")).hexdigest())
app.secret_key = _SECRET_KEY
app.permanent_session_lifetime = timedelta(days=30)

UPLOAD_FOLDER = "uploads"
TODO_UPLOAD_FOLDER = os.path.join(UPLOAD_FOLDER, 'todo')
ALLOWED_EXTENSIONS = {"jpg", "jpeg", "png", "webp", "heic", "gif"}
MAX_UPLOAD_SIZE = 16 * 1024 * 1024
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(TODO_UPLOAD_FOLDER, exist_ok=True)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_SIZE

def run_startup_self_check():
    checks = {"uploads": False, "todo_uploads": False, "backups": False, "todos": False}
    try:
        os.makedirs(UPLOAD_FOLDER, exist_ok=True)
        checks["uploads"] = True
        os.makedirs(TODO_UPLOAD_FOLDER, exist_ok=True)
        checks["todo_uploads"] = True
        os.makedirs("backups", exist_ok=True)
        checks["backups"] = True
    except Exception as e:
        try:
            log_error("startup_self_check_dirs", str(e))
        except Exception:
            pass
    try:
        list_todo_items()
        checks["todos"] = True
    except Exception as e:
        try:
            log_error("startup_self_check_todos", str(e))
        except Exception:
            pass
    return checks

init_db()
STARTUP_CHECKS = run_startup_self_check()

PUBLIC_PATHS = {
    "login", "api_login", "health", "static"
}

def current_username():
    return session.get("user", "")


SYNC_SETTINGS_KEY = 'sync_last_event'
LAST_DAILY_BACKUP_KEY = 'last_daily_backup_date'
PENDING_QUEUE_LIMIT = 50

_db_log_action = log_action

def notify_sync_event(kind='refresh', module='all', message='', extra=None):
    payload = {
        'id': str(int(time.time() * 1000)),
        'kind': kind,
        'module': module or 'all',
        'message': message or '',
        'user': current_username(),
        'at': now(),
        'extra': extra or {},
    }
    try:
        set_setting(SYNC_SETTINGS_KEY, json.dumps(payload, ensure_ascii=False))
    except Exception as e:
        try:
            log_error('notify_sync_event', str(e))
        except Exception:
            pass
    return payload


def log_action(username, action):
    _db_log_action(username, action)
    notify_sync_event(kind='log', module='all', message=action, extra={'username': username})


def ensure_daily_backup():
    try:
        today = datetime.now().strftime('%Y-%m-%d')
        if get_setting(LAST_DAILY_BACKUP_KEY, '') == today:
            return
        result = run_daily_backup()
        if result.get('success'):
            set_setting(LAST_DAILY_BACKUP_KEY, today)
    except Exception as e:
        log_error('ensure_daily_backup', str(e))


def request_key_from_payload(data, endpoint=''):
    key = (request.headers.get('X-Request-Key') or (data or {}).get('request_key') or '').strip()
    # 沒有 request_key 時要照常送出；只有「帶了 request_key 且重複」才擋掉。
    if not key:
        return True
    if register_submit_request(key, endpoint=endpoint):
        return key
    return False

def duplicate_success(message='重複送出已忽略'):
    return jsonify(success=True, duplicate=True, message=message)



def resolve_customer_region(customer_name='', requested_region=''):
    requested = (requested_region or '').strip()
    if requested in ['北區', '中區', '南區']:
        return requested
    if customer_name:
        row = get_customer(customer_name, include_archived=True)
        if row and (row.get('region') or '').strip() in ['北區', '中區', '南區']:
            return (row.get('region') or '').strip()
    return ''


def build_customer_payload_snapshot(customer_name=''):
    customer_name = (customer_name or '').strip()
    customer = get_customer(customer_name, include_archived=True) if customer_name else None
    counts = get_customer_relation_counts(customer_name) if customer_name else {}
    return {'customer': customer, 'relation_counts': counts}


def safe_list_todos(fallback_item=None):
    try:
        return list_todo_items()
    except Exception as e:
        log_error('safe_list_todos', str(e))
        return [fallback_item] if fallback_item else []


def export_rows_to_xlsx(sheet_name, rows, columns):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] or 'Sheet1'
    ws.append([header for header, _ in columns])
    for row in rows:
        ws.append([row.get(key, '') if isinstance(row, dict) else '' for _, key in columns])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(40, max(10, max_len + 2))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def require_login():
    return bool(current_username())

def login_required_json(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not require_login():
            return jsonify(success=False, error="請先登入"), 401
        return f(*args, **kwargs)
    return wrapper

@app.after_request
def add_no_cache_headers(response):
    # 避免手機 / PWA / 瀏覽器一直吃到舊的 HTML、JS、CSS。
    if request.path.startswith('/static/') or not request.path.startswith('/api/'):
        response.headers['Cache-Control'] = 'no-store, no-cache, max-age=0, must-revalidate'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
    return response

@app.before_request
def protect_pages():
    path = request.path
    # FIX52：不要在每次開頁時同步執行備份，避免當天第一個使用者卡住。
    # 需要自動每日備份時，可在 Render 環境變數設定 YX_AUTO_DAILY_BACKUP=1。
    if os.getenv("YX_AUTO_DAILY_BACKUP", "0") == "1" and require_login() and not path.startswith("/static/") and path not in ("/health", "/api/health"):
        ensure_daily_backup()
    if path.startswith("/static/") or path in ("/health",):
        return None
    public = [
        "/login", "/api/login", "/api/health", "/api/native-shell/config",
        "/sw.js", "/manifest.webmanifest"
    ]
    if path in public:
        return None
    if not require_login() and path not in ("/",):
        # Let / redirect to login
        if path.startswith("/api/"):
            return jsonify(success=False, error="請先登入"), 401
        return redirect(url_for("login_page"))
    return None


@app.route("/sw.js")
def serve_root_service_worker():
    resp = send_from_directory(app.static_folder, "service-worker.js", mimetype="application/javascript")
    resp.headers["Cache-Control"] = "no-store, no-cache, max-age=0, must-revalidate"
    resp.headers["Service-Worker-Allowed"] = "/"
    return resp

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

def error_response(msg, code=400):
    return jsonify({"success": False, "error": msg}), code

def compress_image(path):
    try:
        img = Image.open(path)
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        if img.width > 1800:
            ratio = 1800 / float(img.width)
            img = img.resize((1800, int(img.height * ratio)))
        img.save(path, "JPEG", quality=78, optimize=True)
    except Exception as e:
        log_error("compress_image", str(e))

def parse_lines_to_items(text):
    parsed = parse_ocr_text(text)
    return parsed["items"], parsed["text"]



def normalize_item_quantity(product_text, qty=0):
    return effective_product_qty(product_text, qty)


def normalize_item_for_save(item):
    product_text = format_product_text_height2((item.get('product_text') or item.get('product') or '').strip())
    material = clean_material_value(item.get('material') or item.get('product_code') or '', product_text)
    product_code = material
    qty = normalize_item_quantity(product_text, item.get('qty') or 0)
    return {'product_text': product_text, 'product_code': product_code, 'material': material, 'qty': qty}


def aggregate_customer_items(items):
    """Group customer items by source + size + material, show supports/notes, and sort 高 > 寬 > 長 ascending."""
    buckets = {}
    for row in items or []:
        product_text = format_product_text_height2((row.get('product_text') or '').strip())
        if not product_text:
            continue
        source = row.get('source') or ''
        material = (row.get('material') or ((row.get('product_code') or '') if (row.get('product_code') or '') != product_text else '')).strip()
        size = product_display_size(product_text)
        qty = normalize_item_quantity(product_text, row.get('qty') or 0)
        support = product_support_text(product_text)
        # If the right side is only 支數, append x件數 for display. 括號備註會保留。
        if support and ('+' not in support and '＋' not in support and 'x' not in support.lower()):
            support = f"{support}x{qty}"
        elif not support:
            support = str(qty)
        key = (source, size, material)
        if key not in buckets:
            out = dict(row)
            out['qty'] = qty
            out['product_text'] = f"{size}={support}" if support else size
            out['material'] = material
            out['product_code'] = material
            out['size_text'] = size
            out['support_text'] = support
            buckets[key] = out
        else:
            buckets[key]['qty'] = int(buckets[key].get('qty') or 0) + qty
            old_support = (buckets[key].get('support_text') or '').strip()
            if support:
                supports = [x for x in old_support.split('+') if x] if old_support else []
                if support not in supports:
                    supports.append(support)
                buckets[key]['support_text'] = '+'.join(supports)
                buckets[key]['product_text'] = f"{size}={buckets[key]['support_text']}"
    rows = list(buckets.values())
    rows.sort(key=lambda r: (product_sort_tuple(r.get('product_text') or ''), r.get('source') or '', r.get('id') or 0))
    return rows


def warehouse_item_size_key(text):
    raw = str(text or '').replace('×', 'x').replace('X', 'x').replace('＊', 'x').replace('*', 'x').replace('＝', '=').strip()
    left = (raw.split('=', 1)[0].strip() or raw).lower()
    parts = [p for p in left.split('x') if p != '']
    if len(parts) >= 3 and all(part.strip().isdigit() for part in parts[:3]):
        return 'x'.join(str(int(part.strip())) for part in parts[:3])
    return left

def safe_cell_items(cell):
    try:
        return json.loads(cell.get('items_json') or '[]')
    except Exception:
        return []

def warehouse_source_totals():
    totals = {}
    details = {}
    source_rows = []
    for row in list_inventory():
        source_rows.append(('庫存', row))
    for row in get_orders():
        source_rows.append(('訂單', row))
    for row in get_master_orders():
        source_rows.append(('總單', row))
    for source_label, row in source_rows:
        product = (row.get('product_text') or row.get('product') or '').strip()
        size = warehouse_item_size_key(product)
        if not size:
            continue
        customer = (row.get('customer_name') or '').strip()
        try:
            qty = int(row.get('qty') or 0)
        except Exception:
            qty = 0
        if qty <= 0:
            continue
        key = (size, customer)
        totals[key] = totals.get(key, 0) + qty
        details.setdefault(key, []).append({
            'source': source_label,
            'id': row.get('id'),
            'product_text': product,
            'qty': qty,
            'customer_name': customer,
        })
    return totals, details

def warehouse_placed_totals(exclude_cell=None, proposed_items=None):
    placed = {}
    exclude_cell = exclude_cell or None
    for cell in warehouse_get_cells():
        cell_key = (str(cell.get('zone')), int(cell.get('column_index') or 0), int(cell.get('slot_number') or 0))
        if exclude_cell and cell_key == exclude_cell:
            items = proposed_items or []
        else:
            items = safe_cell_items(cell)
        for it in items:
            size = warehouse_item_size_key(it.get('product_text') or it.get('product') or '')
            if not size:
                continue
            customer = (it.get('customer_name') or '').strip()
            try:
                qty = int(it.get('qty') or 0)
            except Exception:
                qty = 0
            if qty <= 0:
                continue
            key = (size, customer)
            placed[key] = placed.get(key, 0) + qty
    return placed

def normalize_warehouse_payload_items(items):
    merged = {}
    for raw in (items or []):
        if not isinstance(raw, dict):
            continue
        product = (raw.get('product_text') or raw.get('product') or '').strip()
        if not product:
            continue
        try:
            qty = int(raw.get('qty') or 0)
        except Exception:
            qty = 0
        if qty <= 0:
            continue
        customer = (raw.get('customer_name') or '').strip()
        key = (warehouse_item_size_key(product), customer)
        if key not in merged:
            item = dict(raw)
            item['product_text'] = product
            item['product_code'] = (raw.get('product_code') or product)
            item['customer_name'] = customer
            item['qty'] = qty
            merged[key] = item
        else:
            merged[key]['qty'] = int(merged[key].get('qty') or 0) + qty
            if not merged[key].get('source_summary') and raw.get('source_summary'):
                merged[key]['source_summary'] = raw.get('source_summary')
    return list(merged.values())

def validate_warehouse_cell_quantities(zone, column_index, slot_number, items):
    """防止從尚未入倉清單拖入超過來源數量。手動新增、來源不存在的商品保留可用。"""
    source_totals, _details = warehouse_source_totals()
    exclude_key = (str(zone), int(column_index), int(slot_number))
    placed = warehouse_placed_totals(exclude_cell=exclude_key, proposed_items=items)
    for it in items:
        size = warehouse_item_size_key(it.get('product_text') or it.get('product') or '')
        customer = (it.get('customer_name') or '').strip()
        key = (size, customer)
        source_total = int(source_totals.get(key, 0) or 0)
        placed_total = int(placed.get(key, 0) or 0)
        from_unplaced = bool(it.get('source_summary') or it.get('source') == 'unplaced')
        if source_total > 0 and placed_total > source_total:
            return False, f"{it.get('product_text') or size} 的入倉數量超過來源數量（來源 {source_total}，目前要放 {placed_total}）"
        if from_unplaced and source_total <= 0:
            return False, f"{it.get('product_text') or size} 找不到可入倉來源，請重新整理後再試"
    return True, ''

def grouped_inventory():
    return inventory_summary()


def resolve_customer_identity(customer_name='', customer_uid='', include_archived=True):
    uid = (customer_uid or '').strip()
    name = (customer_name or '').strip()
    row = None
    if uid:
        row = get_customer_by_uid(uid, include_archived=include_archived)
    if not row and name:
        row = get_customer(name, include_archived=include_archived)
    resolved_name = (row.get('name') if row else name) or ''
    resolved_uid = (row.get('customer_uid') if row else uid) or ''
    return row, resolved_name, resolved_uid


def customer_groups():
    customers = get_customers()
    groups = {"北區": [], "中區": [], "南區": [], "未分區": []}
    for c in customers:
        region = (c.get("region") or '').strip()
        if region not in groups:
            region = "未分區"
        groups[region].append(c)
    return groups

@app.route("/")
def home():
    if not require_login():
        return redirect(url_for("login_page"))
    return render_template("index.html", username=current_username(), title="沅興木業", today=datetime.now().strftime('%Y-%m-%d'))

@app.route("/login")
def login_page():
    if require_login():
        return redirect(url_for("home"))
    return render_template("login.html", title="登入")

@app.route("/settings")
def settings_page():
    is_admin = current_username() == '陳韋廷'
    return render_template("settings.html", username=current_username(), title="設定", is_admin=is_admin, native_ocr_mode=(str(get_setting('native_ocr_mode', '1')) == '1'))

@app.route("/inventory")
def inventory_page():
    return render_template("module.html", module_key="inventory", title="庫存", username=current_username())

@app.route("/orders")
def orders_page():
    return render_template("module.html", module_key="orders", title="訂單", username=current_username())

@app.route("/master-order")
def master_order_page():
    return render_template("module.html", module_key="master_order", title="總單", username=current_username())

@app.route("/ship")
def ship_page():
    return render_template("module.html", module_key="ship", title="出貨", username=current_username())

@app.route("/shipping-query")
def shipping_query_page():
    return render_template("module.html", module_key="shipping_query", title="出貨查詢", username=current_username())

@app.route("/warehouse")
def warehouse_page():
    return render_template("module.html", module_key="warehouse", title="倉庫圖", username=current_username())

@app.route("/customers")
def customers_page():
    return render_template("module.html", module_key="customers", title="客戶資料", username=current_username())

@app.route("/todos")
def todos_page():
    return render_template("module.html", module_key="todos", title="代辦事項", username=current_username())

@app.route("/today-changes")
def today_changes_page():
    return render_template("today_changes.html", username=current_username(), title="今日異動")

@app.route('/todo-image/<path:filename>')
def todo_image(filename):
    if not require_login():
        return redirect(url_for('login_page'))
    safe_name = os.path.basename(filename)
    return send_from_directory(TODO_UPLOAD_FOLDER, safe_name)



@app.route('/api/todos', methods=['GET', 'POST'])
@login_required_json
def api_todos():
    try:
        if request.method == 'GET':
            return jsonify(success=True, items=safe_list_todos())
        files = []
        for key in ('images', 'image'):
            files.extend([f for f in request.files.getlist(key) if f and (f.filename or '').strip()])
        if not files:
            return error_response('請先選擇照片')
        save_names = []
        for file in files:
            if not allowed_file(file.filename):
                return error_response('圖片格式不支援')
            filename = secure_filename(file.filename or '')
            ext = (filename.rsplit('.', 1)[-1].lower() if '.' in filename else 'jpg') or 'jpg'
            save_name = f"todo_{int(time.time()*1000)}_{hashlib.md5((filename+str(time.time())).encode('utf-8')).hexdigest()[:10]}.{ext}"
            save_path = os.path.join(TODO_UPLOAD_FOLDER, save_name)
            file.save(save_path)
            compress_image(save_path)
            save_names.append(save_name)
        note = (request.form.get('note') or '').strip()
        due_date = (request.form.get('due_date') or '').strip()
        created_by = current_username()
        image_payload = json.dumps(save_names, ensure_ascii=False)
        create_todo_item(note=note, due_date=due_date, image_filename=image_payload, created_by=created_by)
        fallback = {'note': note, 'due_date': due_date, 'image_filename': image_payload, 'created_by': created_by, 'created_at': now()}
        log_action(created_by, f"新增代辦 {note or ','.join(save_names)}")
        add_audit_trail(created_by, 'create', 'todo_items', note or 'todo', before_json={}, after_json={'note': note, 'due_date': due_date, 'images': save_names})
        return jsonify(success=True, items=safe_list_todos(fallback_item=fallback))
    except Exception as e:
        log_error('api_todos', str(e))
        return error_response('代辦事項儲存失敗')


@app.route('/api/todos/<int:todo_id>/complete', methods=['POST'])
@login_required_json
def api_todo_complete(todo_id):
    try:
        item = get_todo_item(todo_id)
        if not item:
            return error_response('找不到代辦事項', 404)
        complete_todo_item(todo_id)
        log_action(current_username(), f"完成代辦 {todo_id}")
        return jsonify(success=True, items=safe_list_todos())
    except Exception as e:
        log_error('api_todo_complete', str(e))
        return error_response('代辦事項完成失敗')

@app.route('/api/todos/<int:todo_id>/restore', methods=['POST'])
@login_required_json
def api_todo_restore(todo_id):
    try:
        item = get_todo_item(todo_id)
        if not item:
            return error_response('找不到代辦事項', 404)
        restore_todo_item(todo_id)
        log_action(current_username(), f"還原代辦 {todo_id}")
        return jsonify(success=True, items=safe_list_todos())
    except Exception as e:
        log_error('api_todo_restore', str(e))
        return error_response('代辦事項還原失敗')

@app.route('/api/todos/reorder', methods=['POST'])
@login_required_json
def api_todo_reorder():
    try:
        data = request.get_json(silent=True) or {}
        reorder_todo_items(data.get('ids') or [], done_flag=int(data.get('done_flag') or 0))
        log_action(current_username(), '拖拉排序代辦')
        return jsonify(success=True, items=safe_list_todos())
    except Exception as e:
        log_error('api_todo_reorder', str(e))
        return error_response('代辦排序失敗')

@app.route('/api/todos/<int:todo_id>', methods=['DELETE'])
@login_required_json
def api_todo_delete(todo_id):
    try:
        item = get_todo_item(todo_id)
        if not item:
            return error_response('找不到代辦事項', 404)
        delete_todo_item(todo_id)
        image_raw = item.get('image_filename') or ''
        try:
            image_names = json.loads(image_raw) if str(image_raw).strip().startswith('[') else [image_raw]
        except Exception:
            image_names = [image_raw]
        for image_filename in [os.path.basename(v or '') for v in image_names if v]:
            if image_filename:
                path = os.path.join(TODO_UPLOAD_FOLDER, image_filename)
                if os.path.exists(path):
                    try:
                        os.remove(path)
                    except Exception:
                        pass
        log_action(current_username(), f"刪除代辦 {todo_id}")
        return jsonify(success=True)
    except Exception as e:
        log_error('api_todo_delete', str(e))
        return error_response('代辦事項刪除失敗')

@app.route("/api/login", methods=["POST"])
def api_login():
    try:
        data = request.get_json(silent=True) or {}
        username = (data.get("username") or data.get("name") or "").strip()
        password = (data.get("password") or "").strip()
        if not username or not password:
            return error_response("帳號密碼不可空白")
        user = get_user(username)
        if user and int(user.get('is_blocked') or 0) == 1:
            try:
                log_action(username or 'unknown', '黑名單登入攔截')
            except Exception:
                pass
            return error_response("此帳號已被停用", 403)
        if not user:
            create_user(username, password)
            log_action(username, "建立帳號")
            user = get_user(username) or {}
        else:
            if not verify_password(user.get('password'), password):
                return error_response("密碼錯誤", 403)
            # 舊明碼資料第一次成功登入後自動升級為 hash
            if user.get('password') == password:
                update_password(username, password)
                user = get_user(username) or user
        session.permanent = True
        session["user"] = username
        session["role"] = user.get('role') or ("admin" if username == "陳韋廷" else "user")
        log_action(username, "登入系統")
        return jsonify(success=True, username=username, role=session.get("role"))
    except Exception as e:
        log_error("api_login", str(e))
        return error_response("登入失敗")

@app.route("/api/logout", methods=["POST"])
@login_required_json
def api_logout():
    user = current_username()
    session.clear()
    try:
        log_action(user, "登出系統")
    except Exception:
        pass
    return jsonify(success=True)

@app.route("/api/change_password", methods=["POST"])
@login_required_json
def api_change_password():
    try:
        data = request.get_json(silent=True) or {}
        old_password = (data.get("old_password") or "").strip()
        new_password = (data.get("new_password") or "").strip()
        confirm_password = (data.get("confirm_password") or "").strip()
        user = get_user(current_username())
        if not user or not verify_password(user.get('password'), old_password):
            return error_response("舊密碼錯誤")
        if not new_password or len(new_password) < 4:
            return error_response("新密碼至少 4 碼")
        if new_password != confirm_password:
            return error_response("兩次密碼不一致")
        update_password(current_username(), new_password)
        log_action(current_username(), "修改密碼")
        return jsonify(success=True)
    except Exception as e:
        log_error("change_password", str(e))
        return error_response("修改失敗")

@app.route("/api/native-ocr/parse", methods=["POST"])
@login_required_json
def api_native_ocr_parse():
    try:
        data = request.get_json(silent=True) or {}
        raw_text = (data.get("raw_text") or data.get("text") or "").strip()
        customer_hint = (data.get("customer_hint") or data.get("customer_name") or "").strip()
        native_confidence = int(data.get("confidence") or data.get("ocr_confidence") or 0)
        blocks = data.get("blocks") or data.get("positions") or []
        ocr_mode = (data.get("ocr_mode") or data.get("mode") or 'blue').strip() or 'blue'
        roi = data.get("roi") or None
        if not raw_text and not customer_hint and not blocks:
            return error_response("沒有可解析的辨識文字")
        result = process_native_ocr_text(
            raw_text,
            customer_hint=customer_hint,
            native_confidence=native_confidence,
            blocks=blocks,
            ocr_mode=ocr_mode,
            roi=roi,
        )
        items = result.get('items') or []
        normalized_text = result.get('text') or ''
        customer_guess = result.get('customer_guess') or ''
        partial = bool((normalized_text or raw_text) and (not normalized_text or not customer_guess))
        log_action(current_username(), f"原生OCR辨識[{','.join(result.get('engines', []))}]")
        return jsonify(
            success=True,
            text=normalized_text or raw_text,
            raw_text=result.get('raw_text') or raw_text,
            items=items,
            confidence=int(result.get('confidence') or 0),
            ocr_confidence=int(result.get('ocr_confidence') or native_confidence or 0),
            parse_confidence=int(result.get('parse_confidence') or 0),
            warning=result.get('warning') or '',
            engines=result.get('engines', []),
            customer_guess=customer_guess,
            cleaned_text=result.get('cleaned_text') or '',
            suggested_roi=result.get('suggested_roi'),
            partial=partial,
            line_map=result.get('line_map', []),
            ocr_mode=ocr_mode,
        )
    except Exception as e:
        log_error("native_ocr_parse", str(e))
        return error_response("原生 OCR 文字解析失敗")

@app.route("/api/save_correction", methods=["POST"])
@login_required_json
def api_save_correction():
    try:
        data = request.get_json(silent=True) or {}
        wrong = (data.get("wrong_text") or "").strip()
        correct = (data.get("correct_text") or "").strip()
        if wrong and correct and wrong != correct:
            save_correction(wrong, correct)
            log_action(current_username(), f"修正OCR {wrong}->{correct}")
            add_audit_trail(current_username(), 'upsert', 'corrections', wrong, before_json={}, after_json={'wrong_text': wrong, 'correct_text': correct})
            notify_sync_event(kind='refresh', module='settings', message='OCR 修正詞庫已更新', extra={'wrong_text': wrong})
        return jsonify(success=True)
    except Exception as e:
        log_error("save_correction", str(e))
        return error_response("儲存失敗")

def _parse_items_from_request(data):
    items = data.get("items") or []
    payload_material = (data.get("material") or "").strip().upper()
    if items:
        cleaned = []
        for it in items:
            if payload_material and not (it.get("material") or "").strip():
                it = {**it, "material": payload_material, "product_code": payload_material}
            fixed = normalize_item_for_save(it)
            if int(fixed.get("qty") or 0) <= 0 or not fixed.get("product_text"):
                continue
            cleaned.append(fixed)
        return cleaned
    text = data.get("ocr_text") or data.get("text") or ""
    parsed_items, _ = parse_lines_to_items(text)
    cleaned = []
    for it in parsed_items:
        if payload_material:
            it = {**it, "material": payload_material, "product_code": payload_material}
        fixed = normalize_item_for_save(it)
        if fixed.get("product_text") and int(fixed.get("qty") or 0) > 0:
            cleaned.append(fixed)
    return cleaned

@app.route("/api/inventory", methods=["GET", "POST"])
@login_required_json
def api_inventory():
    try:
        if request.method == "GET":
            return jsonify(success=True, items=grouped_inventory())
        data = request.get_json(silent=True) or {}
        if not request_key_from_payload(data, endpoint='/api/inventory'):
            return duplicate_success('相同庫存送出已忽略')
        items = _parse_items_from_request(data)
        if not items:
            return error_response("請輸入商品資料")
        operator = current_username()
        location = (data.get("location") or "").strip()
        customer_name = (data.get("customer_name") or "").strip()
        if customer_name:
            upsert_customer(customer_name, region=resolve_customer_region(customer_name, data.get('region')))
        for it in items:
            save_inventory_item(it["product_text"], it.get("product_code", ""), int(it["qty"]), location, customer_name, operator, data.get("ocr_text", ""), it.get("material",""))
        log_action(operator, "建立庫存")
        add_audit_trail(operator, 'create', 'inventory', customer_name or 'inventory', before_json={}, after_json={'customer_name': customer_name, 'location': location, 'items': items})
        notify_sync_event(kind='refresh', module='inventory', message='庫存已更新', extra={'customer_name': customer_name, 'count': len(items)})
        snap = build_customer_payload_snapshot(customer_name) if customer_name else {}
        return jsonify(success=True, items=grouped_inventory(), **snap)
    except Exception as e:
        log_error("inventory", str(e))
        return error_response("建立失敗")


@app.route("/api/inventory/<int:item_id>", methods=["GET", "PUT", "DELETE"])
@login_required_json
def api_inventory_item(item_id):
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute(sql("SELECT * FROM inventory WHERE id = ?"), (item_id,))
        row = fetchone_dict(cur)
        if not row:
            conn.close()
            return error_response("找不到庫存商品", 404)
        if request.method == "GET":
            conn.close()
            return jsonify(success=True, item=row)
        if request.method == "DELETE":
            before = dict(row)
            cur.execute(sql("DELETE FROM inventory WHERE id = ?"), (item_id,))
            conn.commit()
            conn.close()
            log_action(current_username(), f"刪除庫存商品 #{item_id}")
            add_audit_trail(current_username(), 'delete', 'inventory', str(item_id), before_json=before, after_json={})
            notify_sync_event(kind='refresh', module='inventory', message='庫存商品已刪除', extra={'id': item_id})
            return jsonify(success=True, items=grouped_inventory())
        data = request.get_json(silent=True) or {}
        product_text = format_product_text_height2((data.get('product_text') or row.get('product_text') or '').strip())
        material = clean_material_value(data.get('material') if data.get('material') is not None else (data.get('product_code') if data.get('product_code') is not None else row.get('material') or row.get('product_code') or ''), product_text)
        product_code = material
        qty = int(data.get('qty') if data.get('qty') is not None else row.get('qty') or 0)
        location = (data.get('location') if data.get('location') is not None else row.get('location') or '').strip()
        customer_name = (data.get('customer_name') if data.get('customer_name') is not None else row.get('customer_name') or '').strip()
        if not product_text:
            conn.close()
            return error_response('請輸入商品資料')
        if qty < 0:
            qty = 0
        before = dict(row)
        cur.execute(sql("""
            UPDATE inventory
            SET product_text = ?, product_code = ?, material = ?, qty = ?, location = ?, customer_name = ?, operator = ?, updated_at = ?
            WHERE id = ?
        """), (product_text, product_code, material, qty, location, customer_name, current_username(), now(), item_id))
        conn.commit()
        conn.close()
        log_action(current_username(), f"編輯庫存商品 #{item_id}")
        add_audit_trail(current_username(), 'update', 'inventory', str(item_id), before_json=before, after_json={'product_text': product_text, 'qty': qty, 'location': location, 'customer_name': customer_name})
        notify_sync_event(kind='refresh', module='inventory', message='庫存商品已更新', extra={'id': item_id})
        return jsonify(success=True, items=grouped_inventory())
    except Exception as e:
        log_error('inventory_item', str(e))
        return error_response('庫存商品操作失敗')

@app.route("/api/inventory/<int:item_id>/move", methods=["POST"])
@login_required_json
def api_inventory_item_move(item_id):
    try:
        data = request.get_json(silent=True) or {}
        target = (data.get('target') or '').strip()
        customer_name = (data.get('customer_name') or '').strip()
        if target not in ('orders', 'master_order', 'master_orders'):
            return error_response('請選擇要移到訂單或總單')
        if not customer_name:
            return error_response('請選擇客戶')
        conn = get_db()
        cur = conn.cursor()
        cur.execute(sql("SELECT * FROM inventory WHERE id = ?"), (item_id,))
        row = fetchone_dict(cur)
        if not row:
            conn.close()
            return error_response('找不到庫存商品', 404)
        current_qty = int(row.get('qty') or 0)
        move_qty = int(data.get('qty') or current_qty or 0)
        if move_qty <= 0:
            conn.close()
            return error_response('移動數量必須大於 0')
        if move_qty > current_qty:
            move_qty = current_qty
        product_text = format_product_text_height2((row.get('product_text') or '').strip())
        product_code = clean_material_value(row.get('material') or row.get('product_code') or '', product_text)
        conn.close()
        upsert_customer(customer_name, region=resolve_customer_region(customer_name, data.get('region')))
        item = {'product_text': product_text, 'product_code': product_code, 'qty': move_qty}
        if target == 'orders':
            save_order(customer_name, [item], current_username(), (data.get('duplicate_mode') or 'merge').strip() or 'merge')
            target_label = '訂單'
            module = 'orders'
        else:
            save_master_order(customer_name, [item], current_username(), (data.get('duplicate_mode') or 'merge').strip() or 'merge')
            target_label = '總單'
            module = 'master_order'
        conn = get_db()
        cur = conn.cursor()
        if move_qty >= current_qty:
            cur.execute(sql("DELETE FROM inventory WHERE id = ?"), (item_id,))
        else:
            cur.execute(sql("UPDATE inventory SET qty = qty - ?, operator = ?, updated_at = ? WHERE id = ?"), (move_qty, current_username(), now(), item_id))
        conn.commit()
        conn.close()
        log_action(current_username(), f"庫存移到{target_label}：{customer_name}")
        add_audit_trail(current_username(), 'move', 'inventory', str(item_id), before_json={'id': item_id, 'qty': current_qty, 'product_text': product_text}, after_json={'target': target_label, 'customer_name': customer_name, 'qty': move_qty, 'product_text': product_text})
        notify_sync_event(kind='refresh', module='inventory', message=f'庫存已移到{target_label}', extra={'id': item_id, 'customer_name': customer_name, 'qty': move_qty})
        notify_sync_event(kind='refresh', module=module, message=f'{target_label}已更新', extra={'customer_name': customer_name, 'qty': move_qty})
        snap = build_customer_payload_snapshot(customer_name)
        return jsonify(success=True, items=grouped_inventory(), customer_name=customer_name, target=target_label, **snap)
    except Exception as e:
        log_error('inventory_item_move', str(e))
        return error_response('庫存移動失敗')
@app.route("/api/orders", methods=["GET", "POST"])
@login_required_json
def api_orders():
    try:
        if request.method == "GET":
            return jsonify(success=True, items=get_orders())
        data = request.get_json(silent=True) or {}
        if not request_key_from_payload(data, endpoint='/api/orders'):
            return duplicate_success('相同訂單送出已忽略')
        items = _parse_items_from_request(data)
        if not items:
            return error_response("請輸入商品資料")
        customer_name = (data.get("customer_name") or "").strip()
        if not customer_name:
            return error_response("請輸入客戶名稱")
        upsert_customer(customer_name, region=resolve_customer_region(customer_name, data.get('region')))
        save_order(customer_name, items, current_username(), (data.get("duplicate_mode") or "merge").strip() or "merge")
        log_action(current_username(), "建立訂單")
        add_audit_trail(current_username(), 'create', 'orders', customer_name, before_json={}, after_json={'customer_name': customer_name, 'items': items})
        notify_sync_event(kind='refresh', module='orders', message='訂單已更新', extra={'customer_name': customer_name, 'count': len(items)})
        snap = build_customer_payload_snapshot(customer_name)
        return jsonify(success=True, items=get_orders(), **snap)
    except Exception as e:
        log_error("orders", str(e))
        return error_response("訂單建立失敗")

@app.route("/api/master_orders", methods=["GET", "POST"])
@login_required_json
def api_master_orders():
    try:
        if request.method == "GET":
            return jsonify(success=True, items=get_master_orders())
        data = request.get_json(silent=True) or {}
        if not request_key_from_payload(data, endpoint='/api/master_orders'):
            return duplicate_success('相同總單送出已忽略')
        items = _parse_items_from_request(data)
        if not items:
            return error_response("請輸入商品資料")
        customer_name = (data.get("customer_name") or "").strip()
        if not customer_name:
            return error_response("請輸入客戶名稱")
        upsert_customer(customer_name, region=resolve_customer_region(customer_name, data.get('region')))
        save_master_order(customer_name, items, current_username(), (data.get("duplicate_mode") or "merge").strip() or "merge")
        log_action(current_username(), "更新總單")
        add_audit_trail(current_username(), 'create', 'master_orders', customer_name, before_json={}, after_json={'customer_name': customer_name, 'items': items})
        notify_sync_event(kind='refresh', module='master_order', message='總單已更新', extra={'customer_name': customer_name, 'count': len(items)})
        snap = build_customer_payload_snapshot(customer_name)
        return jsonify(success=True, items=get_master_orders(), **snap)
    except Exception as e:
        log_error("master_orders", str(e))
        return error_response("總單失敗")

@app.route("/api/ship", methods=["POST"])
@login_required_json
def api_ship():
    try:
        data = request.get_json(silent=True) or {}
        if not request_key_from_payload(data, endpoint='/api/ship'):
            return duplicate_success('相同出貨送出已忽略')
        items = _parse_items_from_request(data)
        if not items:
            return error_response("請輸入商品資料")
        customer_name = (data.get("customer_name") or "").strip()
        if not customer_name:
            return error_response("請輸入客戶名稱")
        upsert_customer(customer_name, region=resolve_customer_region(customer_name, data.get('region')))
        allow_inventory_fallback = bool(data.get("allow_inventory_fallback"))
        result = ship_order(customer_name, items, current_username(), allow_inventory_fallback=allow_inventory_fallback)
        if result.get("success"):
            log_action(current_username(), "完成出貨")
            add_audit_trail(current_username(), 'ship', 'shipping_records', customer_name, before_json={}, after_json={'customer_name': customer_name, 'items': items, 'allow_inventory_fallback': allow_inventory_fallback, 'breakdown': result.get('breakdown', [])})
            notify_sync_event(kind='refresh', module='ship', message='出貨已更新', extra={'customer_name': customer_name, 'count': len(items)})
        if isinstance(result, dict) and customer_name:
            result.update(build_customer_payload_snapshot(customer_name))
        return jsonify(result)
    except Exception as e:
        log_error("ship", str(e))
        return error_response("出貨失敗")

@app.route("/api/shipping_records", methods=["GET"])
@login_required_json
def api_shipping_records():
    start_date = request.args.get("start_date")
    end_date = request.args.get("end_date")
    q = (request.args.get("q") or '').strip()
    rows = get_shipping_records(start_date=start_date, end_date=end_date, q=q)
    return jsonify(success=True, items=rows, records=rows)

@app.route("/api/ship-preview", methods=["POST"])
@login_required_json
def api_ship_preview():
    try:
        data = request.get_json(silent=True) or {}
        items = _parse_items_from_request(data)
        customer_name = (data.get("customer_name") or "").strip()
        if not customer_name:
            return error_response("請輸入客戶名稱")
        if not items:
            return error_response("沒有可預覽的商品")
        return jsonify(preview_ship_order(customer_name, items))
    except Exception as e:
        log_error("ship_preview", str(e))
        return error_response("出貨預覽失敗")

@app.route("/api/customers", methods=["GET", "POST"])
@login_required_json
def api_customers():
    try:
        if request.method == "GET":
            return jsonify(success=True, items=get_customers())
        data = request.get_json(silent=True) or {}
        name = (data.get("name") or "").strip()
        row, resolved_name, _resolved_uid = resolve_customer_identity(name, (data.get('customer_uid') or '').strip(), include_archived=True)
        name = name or resolved_name
        if resolved_name and resolved_name != name and not (data.get('force_new') or False):
            name = resolved_name
        if not name:
            return error_response("請輸入客戶名稱")
        item = upsert_customer(
            name,
            phone=(data.get("phone") or "").strip(),
            address=(data.get("address") or "").strip(),
            notes=(data.get("notes") or "").strip(),
            common_materials=(data.get("common_materials") or "").strip(),
            common_sizes=(data.get("common_sizes") or "").strip(),
            region=resolve_customer_region(name, data.get("region")),
            preserve_existing=bool(data.get('preserve_existing', False))
        )
        log_action(current_username(), f"儲存客戶 {name}")
        add_audit_trail(current_username(), 'upsert', 'customer_profiles', name, before_json=row or {}, after_json=data)
        notify_sync_event(kind='refresh', module='customers', message=f'客戶已更新：{name}', extra={'customer_name': name})
        return jsonify(success=True, items=get_customers(), item=item)
    except Exception as e:
        log_error("customers", str(e))
        return error_response("客戶儲存失敗")


@app.route("/api/customers/archived", methods=["GET"])
@login_required_json
def api_customers_archived():
    try:
        items = [c for c in get_customers(active_only=False) if int(c.get('is_archived') or 0) == 1]
        return jsonify(success=True, items=items)
    except Exception as e:
        log_error("customers_archived", str(e))
        return error_response("封存客戶讀取失敗")

@app.route("/api/customers/<name>/restore", methods=["POST"])
@login_required_json
def api_customer_restore(name):
    try:
        data = request.get_json(silent=True) or {}
        row, resolved_name, _resolved_uid = resolve_customer_identity(name, data.get('customer_uid') or request.args.get('customer_uid') or '', include_archived=True)
        target_name = resolved_name or name
        item = restore_customer(target_name)
        log_action(current_username(), f"復原客戶 {target_name}")
        add_audit_trail(current_username(), 'restore', 'customer_profiles', target_name, before_json={'name': target_name}, after_json={'name': target_name, 'restored': True})
        notify_sync_event(kind='refresh', module='customers', message=f'客戶已復原：{target_name}', extra={'customer_name': target_name})
        return jsonify(success=True, item=item, items=get_customers())
    except Exception as e:
        log_error("restore_customer", str(e))
        return error_response(f"客戶復原失敗：{str(e)}")

@app.route("/api/customers/move", methods=["POST"])
@login_required_json
def api_customers_move():
    try:
        data = request.get_json(silent=True) or {}
        name = (data.get("name") or "").strip()
        region = (data.get("region") or "").strip()
        if region not in ["北區", "中區", "南區"]:
            return error_response("缺少客戶或區域")
        row, resolved_name, _resolved_uid = resolve_customer_identity(name, data.get('customer_uid') or '', include_archived=True)
        name = resolved_name or name
        if not name or not row:
            return error_response("找不到客戶資料")
        before_region = (row.get("region") or "").strip()
        item = upsert_customer(name, phone=(row.get("phone") or "").strip(), address=(row.get("address") or "").strip(), notes=(row.get("notes") or "").strip(), common_materials=(row.get("common_materials") or "").strip(), common_sizes=(row.get("common_sizes") or "").strip(), region=region, preserve_existing=True)
        log_action(current_username(), f"移動客戶 {name} 到 {region}")
        add_audit_trail(current_username(), 'move', 'customer_profiles', name, before_json={'name': name, 'region': before_region}, after_json={'name': name, 'region': region})
        notify_sync_event(kind="refresh", module="customers", message=f"客戶已移動：{name} -> {region}", extra={"customer_name": name, "region": region})
        return jsonify(success=True, items=get_customers(), item=item)
    except Exception as e:
        log_error("move_customer", str(e))
        return error_response("移動客戶失敗")


@app.route("/api/customers/<name>", methods=["GET", "DELETE", "PUT"])
@login_required_json
def api_customer_detail(name):
    if request.method == "PUT":
        try:
            data = request.get_json(silent=True) or {}
            new_name = (data.get("new_name") or "").strip()
            if not new_name:
                return error_response("請輸入新的客戶名稱")
            source, resolved_name, _resolved_uid = resolve_customer_identity(name, (data.get('customer_uid') or '').strip(), include_archived=True)
            name = resolved_name or name
            if not source:
                return error_response("找不到原客戶資料")
            if new_name == name:
                return jsonify(success=True, item=source, counts=get_customer_relation_counts(name))
            existed = get_customer(new_name, include_archived=True)
            if existed:
                return error_response("新的客戶名稱已存在，請換一個名稱")
            conn = get_db()
            cur = conn.cursor()
            try:
                cur.execute(sql("UPDATE customer_profiles SET name = ?, updated_at = ? WHERE name = ?"), (new_name, now(), name))
                cur.execute(sql("UPDATE inventory SET customer_name = ?, customer_uid = ?, updated_at = ? WHERE customer_name = ?"), (new_name, source.get('customer_uid') or '', now(), name))
                cur.execute(sql("UPDATE orders SET customer_name = ?, customer_uid = ?, updated_at = ? WHERE customer_name = ?"), (new_name, source.get('customer_uid') or '', now(), name))
                cur.execute(sql("UPDATE master_orders SET customer_name = ?, customer_uid = ?, updated_at = ? WHERE customer_name = ?"), (new_name, source.get('customer_uid') or '', now(), name))
                cur.execute(sql("UPDATE shipping_records SET customer_name = ?, customer_uid = ? WHERE customer_name = ?"), (new_name, source.get('customer_uid') or '', name))
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                conn.close()
            item = get_customer(new_name, include_archived=True)
            log_action(current_username(), f"修改客戶名稱 {name} -> {new_name}")
            add_audit_trail(current_username(), 'rename', 'customer_profiles', name, before_json={'name': name}, after_json={'name': new_name})
            notify_sync_event(kind="refresh", module="customers", message=f"客戶已改名：{name} -> {new_name}", extra={"customer_name": new_name})
            return jsonify(success=True, item=item, counts=get_customer_relation_counts(new_name))
        except Exception as e:
            log_error("rename_customer", str(e))
            return error_response(f"客戶名稱更新失敗：{str(e)}")
    if request.method == "DELETE":
        try:
            data = request.get_json(silent=True) or {}
            _row, resolved_name, _resolved_uid = resolve_customer_identity(name, data.get('customer_uid') or request.args.get('customer_uid') or '', include_archived=True)
            name = resolved_name or name
            result = delete_customer(name)
            mode = result.get('mode') or 'deleted'
            counts = result.get('counts') or {}
            log_action(current_username(), f"{'封存' if mode == 'archived' else '刪除'}客戶 {name}")
            add_audit_trail(current_username(), 'delete' if mode == 'deleted' else 'archive', 'customer_profiles', name, before_json=result.get('item') or {}, after_json={'mode': mode, 'counts': counts})
            notify_sync_event(kind='refresh', module='customers', message=f"客戶已{'封存' if mode == 'archived' else '刪除'}：{name}", extra={'customer_name': name, 'mode': mode})
            message = '客戶已刪除' if mode == 'deleted' else '客戶已有關聯資料，已改為封存保留歷史資料'
            return jsonify(success=True, mode=mode, counts=counts, message=message)
        except Exception as e:
            log_error("delete_customer", str(e))
            return error_response(f"客戶刪除失敗：{str(e)}")
    row, resolved_name, _resolved_uid = resolve_customer_identity(name, request.args.get('customer_uid') or '', include_archived=True)
    name = resolved_name or name
    if not row:
        return error_response("找不到客戶", 404)
    return jsonify(success=True, item=row, counts=get_customer_relation_counts(name))

@app.route("/api/warehouse", methods=["GET"])
@login_required_json
def api_warehouse():
    try:
        return jsonify(success=True, zones=warehouse_summary(), cells=warehouse_get_cells())
    except Exception as e:
        log_error("api_warehouse", str(e))
        return jsonify(success=True, zones={"A": {}, "B": {}}, cells=[])


@app.route("/api/warehouse/cell", methods=["POST"])
@login_required_json
def api_warehouse_cell():
    try:
        data = request.get_json(silent=True) or {}
        zone = (data.get("zone") or "A").strip().upper()
        column_index = int(data.get("column_index") or 0)
        slot_type = 'direct'
        slot_number = int(data.get("slot_number") or 0)
        if zone not in ("A", "B") or column_index < 1 or column_index > 6 or slot_number < 1:
            return error_response("格位參數錯誤")
        # 防止手動輸入不存在的格位（例如 A-1-99）造成倉庫圖被拉出異常超長格數。
        existing_cells = warehouse_get_cells()
        if not any(str(c.get('zone')) == zone and int(c.get('column_index') or 0) == column_index and int(c.get('slot_number') or 0) == slot_number for c in existing_cells):
            return error_response("格位不存在，請先在格子內點「插入格子」")
        previous_cell = next((c for c in existing_cells if str(c.get('zone')) == zone and int(c.get('column_index') or 0) == column_index and int(c.get('slot_number') or 0) == slot_number), {})
        items = normalize_warehouse_payload_items(data.get("items") or [])
        ok, msg = validate_warehouse_cell_quantities(zone, column_index, slot_number, items)
        if not ok:
            return error_response(msg)
        note = data.get("note") or ""
        warehouse_save_cell(zone, column_index, slot_type, slot_number, items, note)
        if items:
            top_customer = next((it.get('customer_name') for it in items if it.get('customer_name')), '')
            record_recent_slot(current_username(), top_customer, zone, column_index, slot_number)
        log_action(current_username(), f"更新倉庫格位 {zone}{column_index}-{slot_type}-{slot_number}")
        add_audit_trail(current_username(), 'upsert', 'warehouse_cells', f'{zone}-{column_index}-{slot_number}', before_json={'items_json': previous_cell.get('items_json'), 'note': previous_cell.get('note')}, after_json={'zone': zone, 'column_index': column_index, 'slot_number': slot_number, 'items': items, 'note': note})
        notify_sync_event(kind='refresh', module='warehouse', message='倉庫格位已更新', extra={'zone': zone, 'column_index': column_index, 'slot_number': slot_number})
        return jsonify(success=True, zones=warehouse_summary())
    except Exception as e:
        log_error("warehouse_cell", str(e))
        return error_response("格位更新失敗")

@app.route("/api/warehouse/move", methods=["POST"])
@login_required_json
def api_warehouse_move():
    try:
        data = request.get_json(silent=True) or {}
        from_key = data.get("from_key")
        to_key = data.get("to_key")
        product_text = format_product_text_height2(data.get("product_text"))
        customer_name = (data.get("customer_name") or "").strip()
        qty = int(data.get("qty", 1))
        if not (from_key and to_key and product_text):
            return error_response("缺少參數")
        result = warehouse_move_item(tuple(from_key), tuple(to_key), product_text, qty, customer_name=customer_name)
        if result.get("success"):
            log_action(current_username(), f"拖曳商品 {product_text}")
            try:
                to_slot = int(to_key[3] if len(to_key) >= 4 else to_key[2])
                record_recent_slot(current_username(), customer_name, to_key[0], int(to_key[1]), to_slot)
            except Exception:
                pass
            add_audit_trail(current_username(), 'move', 'warehouse_cells', product_text, before_json={'from_key': from_key, 'customer_name': customer_name}, after_json={'to_key': to_key, 'qty': qty, 'product_text': product_text, 'customer_name': customer_name})
            notify_sync_event(kind='refresh', module='warehouse', message='倉庫位置已移動', extra={'product_text': product_text, 'qty': qty, 'customer_name': customer_name})
        return jsonify(result)
    except Exception as e:
        log_error("warehouse_move", str(e))
        return error_response("拖曳失敗")

@app.route("/api/warehouse/add-column", methods=["POST"])
@login_required_json
def api_warehouse_add_column():
    try:
        data = request.get_json(silent=True) or {}
        zone = (data.get("zone") or "A").strip().upper()
        if zone not in ("A", "B"):
            return error_response("區域錯誤")
        column_index = warehouse_add_column(zone)
        log_action(current_username(), f"新增格子欄 {zone}{column_index}")
        notify_sync_event(kind='refresh', module='warehouse', message='倉庫新增欄位', extra={'zone': zone, 'column_index': column_index})
        return jsonify(success=True, column_index=column_index, zones=warehouse_summary(), cells=warehouse_get_cells())
    except Exception as e:
        log_error("warehouse_add_column", str(e))
        return error_response("新增格子失敗")

@app.route("/api/warehouse/search")
@login_required_json
def api_warehouse_search():
    q = (request.args.get("q") or "").strip()
    cells = warehouse_get_cells()
    matched = []
    for cell in cells:
        try:
            items = json.loads(cell.get("items_json") or "[]")
        except Exception:
            items = []
        for it in items:
            hay = f"{cell['zone']} {cell['column_index']} {cell['slot_type']} {cell['slot_number']} {it.get('product_text','')} {it.get('customer_name','')}"
            if not q or q.lower() in hay.lower():
                matched.append({"cell": cell, "item": it})
                break
    return jsonify(success=True, items=matched)

@app.route("/api/warehouse/available-items", methods=["GET"])
@login_required_json
def api_warehouse_available_items():
    """列出尚未放入倉庫圖的商品。

    來源包含：庫存、訂單、總單。相同尺寸會合併；不同客戶分開，避免客戶貨物混在一起。
    """
    try:
        source_totals, source_details = warehouse_source_totals()
        placed = warehouse_placed_totals()
        items = []
        for key, total_qty in source_totals.items():
            size, customer = key
            placed_qty = int(placed.get(key, 0) or 0)
            unplaced_qty = max(0, int(total_qty or 0) - placed_qty)
            if unplaced_qty <= 0:
                continue
            source_qty = {}
            for detail in source_details.get(key, []):
                source_qty[detail['source']] = int(source_qty.get(detail['source'], 0) or 0) + int(detail.get('qty') or 0)
            items.append({
                'product_text': size,
                'product_size': size,
                'customer_name': customer,
                'total_qty': int(total_qty or 0),
                'placed_qty': placed_qty,
                'unplaced_qty': unplaced_qty,
                'qty': unplaced_qty,
                'source_qty': source_qty,
                'sources': [{'source': k, 'qty': v} for k, v in source_qty.items()],
                'source_details': source_details.get(key, []),
                'source_summary': '、'.join([f"{k}{v}" for k, v in source_qty.items()]),
                'needs_red': True,
            })
        items.sort(key=lambda r: (r.get('customer_name') or '未指定客戶', r.get('product_text') or ''))
        return jsonify(success=True, items=items)
    except Exception as e:
        log_error("api_warehouse_available_items", str(e))
        return jsonify(success=True, items=[])


@app.route("/api/customer-items", methods=["GET"])
@login_required_json
def api_customer_items():
    """FIX53：客戶商品直接用 SQL 篩選，不再整表載入後 Python 過濾。"""
    name = (request.args.get("name") or "").strip()
    uid = (request.args.get("customer_uid") or "").strip()
    row, resolved_name, resolved_uid = resolve_customer_identity(name, uid, include_archived=True)
    name = resolved_name or name
    uid = resolved_uid or uid or ((row or {}).get('customer_uid') or '')
    items = []
    if not name and not uid:
        return jsonify(success=True, items=[])

    conn = get_db()
    cur = conn.cursor()
    try:
        def pull(table, source_label):
            if uid and name:
                cur.execute(sql(f"""
                    SELECT * FROM {table}
                    WHERE customer_uid = ? OR (COALESCE(customer_uid, '') = '' AND customer_name = ?)
                    ORDER BY id DESC
                """), (uid, name))
            elif uid:
                cur.execute(sql(f"SELECT * FROM {table} WHERE customer_uid = ? ORDER BY id DESC"), (uid,))
            else:
                cur.execute(sql(f"SELECT * FROM {table} WHERE customer_name = ? ORDER BY id DESC"), (name,))
            for r in rows_to_dict(cur):
                r['source'] = source_label
                items.append(r)
        pull('orders', '訂單')
        pull('master_orders', '總單')
        pull('inventory', '庫存')
    finally:
        conn.close()
    return jsonify(success=True, items=aggregate_customer_items(items))


@app.route("/api/customer-item", methods=["POST", "DELETE"])
@login_required_json
def api_customer_item_modify():
    try:
        data = request.get_json(silent=True) or {}
        source = (data.get("source") or "").strip()
        item_id = int(data.get("id") or 0)
        if not source or not item_id:
            return error_response("缺少商品參數")
        if request.method == "DELETE":
            delete_customer_item(source, item_id)
            log_action(current_username(), f"刪除客戶商品 {source}#{item_id}")
            notify_sync_event(kind='refresh', module='customers', message='客戶商品已刪除', extra={'source': source, 'id': item_id})
            return jsonify(success=True)
        product_text = format_product_text_height2((data.get("product_text") or "").strip())
        qty = int(data.get("qty") or 0)
        if not product_text:
            return error_response("請輸入商品資料")
        material = data.get("material") if "material" in data else None
        update_customer_item(source, item_id, product_text, qty, current_username(), material=material)
        log_action(current_username(), f"更新客戶商品 {source}#{item_id}")
        notify_sync_event(kind='refresh', module='customers', message='客戶商品已更新', extra={'source': source, 'id': item_id})
        return jsonify(success=True)
    except Exception as e:
        log_error("customer_item_modify", str(e))
        return error_response("客戶商品修改失敗")


@app.route("/api/customer-items/batch-material", methods=["POST"])
@login_required_json
def api_customer_items_batch_material():
    try:
        data = request.get_json(silent=True) or {}
        material = (data.get("material") or "").strip().upper()
        items = data.get("items") or []
        if not material:
            return error_response("請選擇材質")
        if not items:
            return error_response("請先勾選要套用材質的商品")
        count = update_items_material(items, material, current_username())
        add_audit_trail(current_username(), 'update', 'customer_items', 'batch_material', before_json={}, after_json={'material': material, 'count': count, 'items': items})
        log_action(current_username(), f"批量套用材質 {material}，共 {count} 筆")
        notify_sync_event(kind='refresh', module='all', message='材質已批量更新', extra={'material': material, 'count': count})
        return jsonify(success=True, count=count, material=material)
    except Exception as e:
        log_error("customer_items_batch_material", str(e))
        return error_response(str(e) or "批量材質更新失敗")



@app.route("/api/customer-items/batch-delete", methods=["POST"])
@login_required_json
def api_customer_items_batch_delete():
    """FIX56：庫存 / 訂單 / 總單共用批量刪除。"""
    try:
        data = request.get_json(silent=True) or {}
        items = data.get("items") or []
        if not items:
            return error_response("請先勾選要刪除的商品")
        table_map = {
            "庫存": "inventory", "inventory": "inventory",
            "訂單": "orders", "orders": "orders",
            "總單": "master_orders", "master_order": "master_orders", "master_orders": "master_orders",
        }
        conn = get_db()
        cur = conn.cursor()
        deleted = 0
        before_items = []
        try:
            for it in items:
                source = (it.get("source") or "").strip()
                table = table_map.get(source)
                item_id = int(it.get("id") or 0)
                if not table or item_id <= 0:
                    continue
                cur.execute(sql(f"SELECT * FROM {table} WHERE id = ?"), (item_id,))
                row = fetchone_dict(cur)
                if not row:
                    continue
                before_items.append({"source": source, "table": table, "id": item_id, "row": row})
                cur.execute(sql(f"DELETE FROM {table} WHERE id = ?"), (item_id,))
                deleted += cur.rowcount or 0
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()
        add_audit_trail(current_username(), "delete", "customer_items", "batch_delete", before_json=before_items, after_json={"count": deleted})
        log_action(current_username(), f"批量刪除商品，共 {deleted} 筆")
        notify_sync_event(kind="refresh", module="all", message="商品已批量刪除", extra={"count": deleted})
        return jsonify(success=True, count=deleted)
    except Exception as e:
        log_error("customer_items_batch_delete", str(e))
        return error_response(str(e) or "批量刪除失敗")

@app.route("/api/customer-items/merge", methods=["POST"])
@login_required_json
def api_customer_items_merge():
    """FIX76：合併同來源、同客戶/位置、同商品的重複資料。"""
    try:
        data = request.get_json(silent=True) or {}
        source = (data.get("source") or "").strip()
        ids = []
        for v in (data.get("ids") or []):
            try:
                iv = int(v)
                if iv > 0 and iv not in ids:
                    ids.append(iv)
            except Exception:
                pass
        table_map = {
            "庫存": "inventory", "inventory": "inventory",
            "訂單": "orders", "orders": "orders",
            "總單": "master_orders", "master_order": "master_orders", "master_orders": "master_orders",
        }
        table = table_map.get(source)
        if not table or len(ids) < 2:
            return error_response("合併參數不足")
        placeholders = ",".join(["?"] * len(ids))
        conn = get_db()
        cur = conn.cursor()
        try:
            cur.execute(sql(f"SELECT * FROM {table} WHERE id IN ({placeholders}) ORDER BY id ASC"), tuple(ids))
            rows = rows_to_dict(cur)
            if len(rows) < 2:
                return error_response("找不到可合併商品")
            keep = rows[0]
            keep_id = int(keep.get("id") or 0)
            total_qty = sum(int(r.get("qty") or 0) for r in rows)
            material = clean_material_value(keep.get("material") or keep.get("product_code") or "", keep.get("product_text") or "")
            product_text = format_product_text_height2(keep.get("product_text") or "")
            cur.execute(sql(f"UPDATE {table} SET qty = ?, product_text = ?, product_code = ?, material = ?, operator = ?, updated_at = ? WHERE id = ?"),
                        (total_qty, product_text, material, material, current_username(), now(), keep_id))
            delete_ids = [int(r.get("id") or 0) for r in rows[1:] if int(r.get("id") or 0) != keep_id]
            if delete_ids:
                delete_ph = ",".join(["?"] * len(delete_ids))
                cur.execute(sql(f"DELETE FROM {table} WHERE id IN ({delete_ph})"), tuple(delete_ids))
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()
        add_audit_trail(current_username(), "merge", table, str(keep_id), before_json={"rows": rows}, after_json={"keep_id": keep_id, "deleted_ids": delete_ids, "qty": total_qty})
        log_action(current_username(), f"合併重複商品 {table} #{keep_id}，共 {len(rows)} 筆")
        notify_sync_event(kind="refresh", module="all", message="重複商品已合併", extra={"source": source, "keep_id": keep_id, "count": len(rows)})
        if table == "inventory":
            return jsonify(success=True, items=grouped_inventory(), keep_id=keep_id, deleted_ids=delete_ids, qty=total_qty)
        if table == "orders":
            return jsonify(success=True, items=get_orders(), keep_id=keep_id, deleted_ids=delete_ids, qty=total_qty)
        return jsonify(success=True, items=get_master_orders(), keep_id=keep_id, deleted_ids=delete_ids, qty=total_qty)
    except Exception as e:
        log_error("customer_items_merge", str(e))
        return error_response(str(e) or "合併商品失敗")

@app.route("/api/backup", methods=["POST", "GET"])
@login_required_json
def api_backup():
    return jsonify(run_daily_backup())

@app.route("/api/backups", methods=["GET"])
@login_required_json
def api_backups():
    return jsonify(list_backups())


@app.route("/api/admin/users", methods=["GET"])
@login_required_json
def api_admin_users():
    if current_username() != '陳韋廷':
        return error_response("權限不足", 403)
    return jsonify(success=True, items=list_users())

@app.route("/api/admin/block", methods=["POST"])
@login_required_json
def api_admin_block():
    if current_username() != '陳韋廷':
        return error_response("權限不足", 403)
    data = request.get_json(silent=True) or {}
    username = (data.get('username') or '').strip()
    blocked = bool(data.get('blocked'))
    if not username or username == '陳韋廷':
        return error_response("不可操作此帳號")
    set_user_blocked(username, blocked)
    log_action(current_username(), f"{'封鎖' if blocked else '解除封鎖'}帳號 {username}")
    notify_sync_event(kind='refresh', module='settings', message='帳號黑名單已更新', extra={'username': username, 'blocked': blocked})
    return jsonify(success=True, items=list_users())



@app.route("/api/warehouse/return-unplaced", methods=["POST"])
@login_required_json
def api_warehouse_return_unplaced():
    """FIX75：把某格已放入的商品清回未錄入倉庫圖狀態。

    倉庫圖的「未錄入」是由來源總量 - 已放入格位數量即時計算，
    所以這裡只要清空該格商品，商品就會自動回到尚未添加倉庫圖清單。
    """
    try:
        data = request.get_json(silent=True) or {}
        zone = (data.get("zone") or "A").strip().upper()
        column_index = int(data.get("column_index") or 0)
        slot_number = int(data.get("slot_number") or 0)
        if zone not in ("A", "B") or column_index < 1 or column_index > 6 or slot_number < 1:
            return error_response("格位參數錯誤")
        cells = warehouse_get_cells()
        cell = next((c for c in cells if str(c.get('zone')) == zone and int(c.get('column_index') or 0) == column_index and int(c.get('slot_number') or 0) == slot_number), None)
        if not cell:
            return error_response("找不到格位")
        items = safe_cell_items(cell)
        note = cell.get('note') or ''
        warehouse_save_cell(zone, column_index, 'direct', slot_number, [], note)
        log_action(current_username(), f"倉庫格位返回上一步 {zone}{column_index}-{slot_number}")
        add_audit_trail(current_username(), 'undo', 'warehouse_cells', f'{zone}-{column_index}-{slot_number}', before_json={'items': items, 'note': note}, after_json={'items': [], 'note': note, 'returned_to_unplaced': True})
        notify_sync_event(kind='refresh', module='warehouse', message='格位商品已回到未錄入倉庫圖', extra={'zone': zone, 'column_index': column_index, 'slot_number': slot_number, 'count': len(items)})
        return jsonify(success=True, returned_items=items, zones=warehouse_summary(), cells=warehouse_get_cells())
    except Exception as e:
        log_error("warehouse_return_unplaced", str(e))
        return error_response("返回上一步失敗")

@app.route("/api/warehouse/add-slot", methods=["POST"])
@login_required_json
def api_warehouse_add_slot():
    try:
        data = request.get_json(silent=True) or {}
        zone = (data.get("zone") or "A").strip().upper()
        column_index = int(data.get("column_index") or 0)
        if zone not in ("A", "B") or column_index < 1 or column_index > 6:
            return error_response("格位參數錯誤")
        slot_type = 'direct'
        insert_after = data.get("insert_after", None)
        if insert_after is None and data.get("slot_number") not in (None, ""):
            insert_after = max(0, int(data.get("slot_number")) - 1)
        slot_number = warehouse_add_slot(zone, column_index, slot_type, insert_after=insert_after)
        log_action(current_username(), f"新增格子 {zone}{column_index}-{slot_number}")
        notify_sync_event(kind='refresh', module='warehouse', message='倉庫新增格子', extra={'zone': zone, 'column_index': column_index, 'slot_number': slot_number, 'insert_after': insert_after})
        return jsonify(success=True, slot_number=slot_number, zones=warehouse_summary(), cells=warehouse_get_cells())
    except Exception as e:
        log_error("warehouse_add_slot", str(e))
        return error_response("新增格子失敗")

@app.route("/api/warehouse/remove-slot", methods=["POST"])
@login_required_json
def api_warehouse_remove_slot():
    try:
        data = request.get_json(silent=True) or {}
        zone = (data.get("zone") or "A").strip().upper()
        column_index = int(data.get("column_index") or 0)
        slot_number = int(data.get("slot_number") or 0)
        if zone not in ("A", "B") or column_index < 1 or column_index > 6 or slot_number < 1:
            return error_response("格位參數錯誤")
        slot_type = 'direct'
        result = warehouse_remove_slot(zone, column_index, slot_type, slot_number)
        if not result.get('success'):
            return error_response(result.get('error') or '刪除格子失敗')
        log_action(current_username(), f"刪除格子 {zone}{column_index}-{slot_number}")
        notify_sync_event(kind='refresh', module='warehouse', message='倉庫刪除格子', extra={'zone': zone, 'column_index': column_index, 'slot_number': slot_number})
        return jsonify(success=True, zones=warehouse_summary(), cells=warehouse_get_cells())
    except Exception as e:
        log_error("warehouse_remove_slot", str(e))
        return error_response("刪除格子失敗")


@app.route("/api/orders/to-master", methods=["POST"])
@login_required_json
def api_orders_to_master():
    try:
        data = request.get_json(silent=True) or {}
        customer_name = (data.get("customer_name") or "").strip()
        product_text = format_product_text_height2((data.get("product_text") or "").strip())
        product_code = (data.get("product_code") or "").strip()
        qty = int(data.get("qty") or 0)
        if not customer_name or not product_text or qty <= 0:
            return error_response("參數不足")
        upsert_customer(customer_name, region=resolve_customer_region(customer_name, data.get('region')))
        save_master_order(customer_name, [{"product_text": product_text, "product_code": product_code, "qty": qty}], current_username())
        log_action(current_username(), f"訂單加入總單 {customer_name} {product_text}x{qty}")
        notify_sync_event(kind='refresh', module='master_order', message='訂單已加入總單', extra={'customer_name': customer_name, 'product_text': product_text, 'qty': qty})
        return jsonify(success=True, items=get_master_orders())
    except Exception as e:
        log_error("orders_to_master", str(e))
        return error_response("加入總單失敗")


def _today_key():
    return datetime.now().strftime('%Y-%m-%d')

def _aggregate_customer_products(rows):
    out = {}
    for r in rows:
        key = ((r.get('customer_name') or '').strip(), (r.get('product_text') or '').strip())
        out[key] = out.get(key, 0) + int(r.get('qty') or 0)
    return out

def _aggregate_inventory_products(rows):
    out = {}
    for r in rows:
        key = (r.get('product_text') or '').strip()
        out[key] = out.get(key, 0) + int(r.get('qty') or 0)
    return out


def _build_anomalies(inv_rows, order_rows, master_rows):
    anomalies = {
        'negative_inventory': [],
        'orders_over_master': [],
        'master_over_inventory': [],
        'unplaced': [],
        'duplicate_products': [],
        'shipping_deduction': [],
        'ocr_errors': [],
        'blocked_logins': [],
    }
    inv_by_product = _aggregate_inventory_products(inv_rows)
    ord_by_cp = _aggregate_customer_products(order_rows)
    mst_by_cp = _aggregate_customer_products(master_rows)

    for r in inv_rows:
        if int(r.get('qty') or 0) < 0:
            anomalies['negative_inventory'].append({'type':'negative_inventory','message':f"庫存負數：{r.get('product_text')} ({r.get('qty')})",'product_text':r.get('product_text') or ''})
        if int(r.get('unplaced_qty') or 0) > 0:
            anomalies['unplaced'].append({'type':'unplaced','message':f"未錄入倉庫圖：{r.get('product_text')} ({r.get('unplaced_qty')})",'product_text':r.get('product_text') or '', 'qty':int(r.get('unplaced_qty') or 0)})

    for key, oq in ord_by_cp.items():
        mq = int(mst_by_cp.get(key) or 0)
        if oq > mq:
            customer_name, product_text = key
            anomalies['orders_over_master'].append({'type':'orders_over_master','message':f"{customer_name}｜訂單大於總單：{product_text} ({oq}>{mq})",'customer_name':customer_name,'product_text':product_text,'order_qty':oq,'master_qty':mq})

    product_master_total = {}
    for (_, product_text), qty in mst_by_cp.items():
        product_master_total[product_text] = product_master_total.get(product_text, 0) + int(qty or 0)
    for product_text, mq in product_master_total.items():
        iq = int(inv_by_product.get(product_text) or 0)
        if mq > iq:
            anomalies['master_over_inventory'].append({'type':'master_over_inventory','message':f"總單大於庫存：{product_text} ({mq}>{iq})",'product_text':product_text,'master_qty':mq,'inventory_qty':iq})

    def _dups(rows, source_name):
        seen = {}
        for r in rows:
            key = ((r.get('customer_name') or '').strip(), (r.get('product_text') or '').strip())
            if not all(key):
                continue
            seen.setdefault(key, []).append(r)
        for key, vals in seen.items():
            if len(vals) > 1:
                anomalies['duplicate_products'].append({'type':'duplicate_products','message':f"{source_name}重複商品：{key[0]}｜{key[1]}（{len(vals)}筆）",'customer_name':key[0],'product_text':key[1],'source':source_name})
    _dups(order_rows, '訂單')
    _dups(master_rows, '總單')

    try:
        conn = get_db(); cur = conn.cursor(); today = _today_key()
        cur.execute(sql("SELECT customer_name, product_text, note, shipped_at FROM shipping_records WHERE substr(shipped_at,1,10)=? ORDER BY shipped_at DESC"), (today,))
        for r in rows_to_dict(cur):
            note = r.get('note') or ''
            if '補扣' in note:
                anomalies['shipping_deduction'].append({'type':'shipping_deduction','message':f"出貨補扣：{r.get('customer_name')}｜{r.get('product_text')}｜{note}"})
        cur.execute(sql("SELECT source, message, created_at FROM errors WHERE substr(created_at,1,10)=? ORDER BY created_at DESC LIMIT 50"), (today,))
        for r in rows_to_dict(cur):
            src = (r.get('source') or '')
            if 'ocr' in src.lower():
                anomalies['ocr_errors'].append({'type':'ocr_errors','message':f"OCR異常：{src}｜{r.get('message') or ''}"})
        cur.execute(sql("SELECT username, action, created_at FROM logs WHERE substr(created_at,1,10)=? AND action LIKE ? ORDER BY created_at DESC LIMIT 50"), (today, '%黑名單登入攔截%'))
        for r in rows_to_dict(cur):
            anomalies['blocked_logins'].append({'type':'blocked_logins','message':f"黑名單登入異常：{r.get('username') or ''}"})
        conn.close()
    except Exception:
        pass

    return anomalies


def _today_changes_payload():
    conn = get_db()
    cur = conn.cursor()
    today = _today_key()
    cur.execute(sql("SELECT id, username, action, created_at FROM logs WHERE substr(created_at,1,10)=? ORDER BY created_at DESC LIMIT 200"), (today,))
    logs = rows_to_dict(cur)
    conn.close()

    inbound = []
    outbound = []
    new_orders = []
    others = []
    for r in logs:
        action = r.get('action') or ''
        if '出貨' in action:
            outbound.append(r)
        elif '訂單' in action:
            new_orders.append(r)
        elif any(k in action for k in ['庫存', '進貨', 'OCR', '總單', '儲存客戶']):
            inbound.append(r)
        else:
            others.append(r)

    inv = inventory_summary()
    orders = get_orders()
    masters = get_master_orders()
    anomalies = _build_anomalies(inv, orders, masters)
    anomaly_list = []
    for key in ['negative_inventory', 'orders_over_master', 'master_over_inventory', 'duplicate_products', 'shipping_deduction', 'ocr_errors', 'blocked_logins']:
        anomaly_list.extend(anomalies.get(key, []))
    unplaced = anomalies['unplaced']
    read_at = get_setting('today_changes_read_at', '') or ''
    unread_count = len([r for r in logs if not read_at or (r.get('created_at') or '') > read_at])

    return {
        'summary': {
            'inbound_count': len(inbound),
            'outbound_count': len(outbound),
            'new_order_count': len(new_orders),
            'unplaced_count': len(unplaced),
            'anomaly_count': len(anomaly_list),
            'unread_count': unread_count,
        },
        'feed': {
            'inbound': inbound[:60],
            'outbound': outbound[:60],
            'new_orders': new_orders[:60],
            'others': others[:40],
        },
        'unplaced_items': unplaced[:120],
        'anomalies': anomaly_list[:120],
        'anomaly_groups': anomalies,
        'read_at': read_at,
    }

@app.route('/api/today-changes', methods=['GET'])
@login_required_json
def api_today_changes():
    return jsonify(success=True, **_today_changes_payload())

@app.route('/api/today-changes/read', methods=['POST'])
@login_required_json
def api_today_changes_mark_read():
    try:
        set_setting('today_changes_read_at', now())
        notify_sync_event(kind='refresh', module='today_changes', message='今日異動已讀已更新')
        return jsonify(success=True)
    except Exception as e:
        log_error('today_changes_mark_read', str(e))
        return error_response('清除已讀失敗')

@app.route('/api/today-changes/<int:log_id>', methods=['DELETE'])
@login_required_json
def api_today_change_delete(log_id):
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute(sql('DELETE FROM logs WHERE id = ?'), (log_id,))
        conn.commit()
        conn.close()
        notify_sync_event(kind='refresh', module='today_changes', message='今日異動已刪除', extra={'log_id': log_id})
        return jsonify(success=True, **_today_changes_payload())
    except Exception as e:
        log_error('today_change_delete', str(e))
        return error_response('刪除異動失敗')

@app.route('/api/anomalies', methods=['GET'])
@login_required_json
def api_anomalies():
    payload = _today_changes_payload()
    return jsonify(success=True, groups=payload.get('anomaly_groups', {}), items=payload.get('anomalies', []), unplaced_items=payload.get('unplaced_items', []))


@app.route('/api/sync/stream')
@login_required_json
def api_sync_stream():
    def generate():
        last_seen = ''
        while True:
            payload = get_setting(SYNC_SETTINGS_KEY, '') or ''
            if payload and payload != last_seen:
                last_seen = payload
                yield f"data: {payload}\n\n"
            else:
                yield ': keepalive\n\n'
            time.sleep(1.5)
    return Response(stream_with_context(generate()), mimetype='text/event-stream', headers={'Cache-Control': 'no-cache', 'X-Accel-Buffering': 'no'})

@app.route('/api/corrections', methods=['GET', 'POST', 'DELETE'])
@login_required_json
def api_corrections_manage():
    try:
        if request.method == 'GET':
            return jsonify(success=True, items=list_corrections_rows())
        data = request.get_json(silent=True) or {}
        wrong_text = (data.get('wrong_text') or '').strip()
        if request.method == 'DELETE':
            delete_correction(wrong_text)
            notify_sync_event(kind='refresh', module='settings', message='OCR 修正詞已刪除', extra={'wrong_text': wrong_text})
            return jsonify(success=True, items=list_corrections_rows())
        correct_text = (data.get('correct_text') or '').strip()
        if not wrong_text or not correct_text:
            return error_response('請輸入錯字與正確字')
        save_correction(wrong_text, correct_text)
        add_audit_trail(current_username(), 'upsert', 'corrections', wrong_text, before_json={}, after_json={'wrong_text': wrong_text, 'correct_text': correct_text})
        notify_sync_event(kind='refresh', module='settings', message='OCR 修正詞已更新', extra={'wrong_text': wrong_text})
        return jsonify(success=True, items=list_corrections_rows())
    except Exception as e:
        log_error('api_corrections_manage', str(e))
        return error_response('修正詞管理失敗')

@app.route('/api/customer-aliases', methods=['GET', 'POST', 'DELETE'])
@login_required_json
def api_customer_aliases_manage():
    try:
        if request.method == 'GET':
            return jsonify(success=True, items=list_customer_aliases())
        data = request.get_json(silent=True) or {}
        alias = (data.get('alias') or '').strip()
        if request.method == 'DELETE':
            delete_customer_alias(alias)
            notify_sync_event(kind='refresh', module='settings', message='客戶別名已刪除', extra={'alias': alias})
            return jsonify(success=True, items=list_customer_aliases())
        target_name = (data.get('target_name') or '').strip()
        if not alias or not target_name:
            return error_response('請輸入別名與正式客戶名稱')
        save_customer_alias(alias, target_name)
        add_audit_trail(current_username(), 'upsert', 'customer_aliases', alias, before_json={}, after_json={'alias': alias, 'target_name': target_name})
        notify_sync_event(kind='refresh', module='settings', message='客戶別名已更新', extra={'alias': alias})
        return jsonify(success=True, items=list_customer_aliases())
    except Exception as e:
        log_error('api_customer_aliases_manage', str(e))
        return error_response('客戶別名管理失敗')

@app.route('/api/recent-slots', methods=['GET'])
@login_required_json
def api_recent_slots():
    customer_name = (request.args.get('customer_name') or '').strip()
    return jsonify(success=True, items=get_recent_slots(current_username(), customer_name=customer_name, limit=8))


# ==== FIX30：操作紀錄中文化 / 陳韋廷權限 / 批量刪除 ====
def _is_admin_user():
    return current_username() == '陳韋廷'

def _parse_maybe_json(value):
    if isinstance(value, (dict, list)):
        return value
    if not value:
        return {}
    try:
        return json.loads(value)
    except Exception:
        return value

def _audit_action_label(action_type=''):
    action_type = (action_type or '').strip()
    return {
        'create': '新增', 'update': '修改', 'delete': '刪除', 'move': '搬移',
        'ship': '出貨', 'transfer': '互通移動', 'upsert': '儲存 / 更新',
        'undo': '還原', 'restore': '還原', 'archive': '封存'
    }.get(action_type, action_type or '操作')

def _audit_entity_label(entity_type=''):
    entity_type = (entity_type or '').strip()
    return {
        'inventory': '庫存', 'orders': '訂單', 'master_orders': '總單',
        'shipping_records': '出貨紀錄', 'warehouse_cells': '倉庫格位',
        'customer_profiles': '客戶資料', 'customer_aliases': '客戶別名',
        'corrections': 'OCR修正詞', 'todo_items': '代辦事項', 'undo': '還原紀錄'
    }.get(entity_type, entity_type or '資料')

def _audit_field_label(key=''):
    return {
        'customer_name': '客戶', 'name': '客戶名稱', 'new_name': '新客戶名稱',
        'product_text': '商品資料', 'product_code': '材質 / 代碼', 'qty': '數量',
        'quantity': '數量', 'location': '倉庫位置', 'operator': '操作人',
        'target': '目標', 'source': '來源', 'source_label': '來源', 'target_label': '目標',
        'zone': '區域', 'column_index': '欄位', 'slot_number': '格號',
        'from_key': '原格位', 'to_key': '新格位', 'note': '備註',
        'items': '商品', 'breakdown': '出貨扣除明細', 'message': '訊息',
        'region': '區域', 'phone': '電話', 'address': '地址', 'notes': '特殊要求',
        'wrong_text': '錯誤字', 'correct_text': '正確字', 'alias': '別名', 'target_name': '正式客戶'
    }.get(key, key)

def _fmt_audit_value(value):
    if isinstance(value, list):
        parts = []
        for item in value:
            if isinstance(item, dict):
                parts.append(_audit_dict_to_text(item).replace('\n', '，'))
            else:
                parts.append(str(item))
        return '；'.join([p for p in parts if p]) or '無'
    if isinstance(value, dict):
        return _audit_dict_to_text(value).replace('\n', '，') or '無'
    if value is None or value == '':
        return '無'
    return str(value)

def _audit_dict_to_text(data):
    data = _parse_maybe_json(data)
    if not data:
        return '無資料'
    if isinstance(data, list):
        return _fmt_audit_value(data)
    if not isinstance(data, dict):
        return str(data)
    preferred = ['source', 'target', 'customer_name', 'name', 'product_text', 'qty', 'location', 'zone', 'column_index', 'slot_number', 'from_key', 'to_key', 'message', 'note']
    lines = []
    used = set()
    for key in preferred:
        if key in data:
            lines.append(f"{_audit_field_label(key)}：{_fmt_audit_value(data.get(key))}")
            used.add(key)
    for key, value in data.items():
        if key in used or key in ('id', 'created_at', 'updated_at', 'customer_uid'):
            continue
        lines.append(f"{_audit_field_label(key)}：{_fmt_audit_value(value)}")
    return '\n'.join(lines) if lines else '無資料'

def _audit_summary_text(item):
    action = _audit_action_label(item.get('action_type'))
    entity = _audit_entity_label(item.get('entity_type'))
    after = _parse_maybe_json(item.get('after_json'))
    before = _parse_maybe_json(item.get('before_json'))
    data = after if isinstance(after, dict) and after else before if isinstance(before, dict) else {}
    customer = (data.get('customer_name') or data.get('name') or '') if isinstance(data, dict) else ''
    product = (data.get('product_text') or '') if isinstance(data, dict) else ''
    qty = (data.get('qty') or data.get('quantity') or '') if isinstance(data, dict) else ''
    target = (data.get('target') or '') if isinstance(data, dict) else ''
    bits = [f"{action}{entity}"]
    if target:
        bits.append(f"目標：{target}")
    if customer:
        bits.append(f"客戶：{customer}")
    if product:
        bits.append(f"商品：{product}")
    if qty != '':
        bits.append(f"數量：{qty}")
    if len(bits) == 1 and item.get('entity_key'):
        bits.append(f"編號：{item.get('entity_key')}")
    return '｜'.join(bits)

def _decorate_audit_item(item):
    before = _parse_maybe_json(item.get('before_json'))
    after = _parse_maybe_json(item.get('after_json'))
    out = dict(item)
    out['action_label'] = _audit_action_label(item.get('action_type'))
    out['entity_label'] = _audit_entity_label(item.get('entity_type'))
    out['summary_text'] = _audit_summary_text(item)
    out['before_text'] = _audit_dict_to_text(before)
    out['after_text'] = _audit_dict_to_text(after)
    return out

def _delete_rows_by_ids(table, ids):
    ids = [int(x) for x in (ids or []) if str(x).isdigit()]
    if not ids:
        return 0
    conn = get_db(); cur = conn.cursor()
    holders = ','.join(['?'] * len(ids))
    cur.execute(sql(f'DELETE FROM {table} WHERE id IN ({holders})'), tuple(ids))
    count = cur.rowcount if cur.rowcount is not None else len(ids)
    conn.commit(); conn.close()
    return count

@app.route('/api/audit-trails', methods=['GET'])
@login_required_json
def api_audit_trails():
    if not _is_admin_user():
        return error_response('操作紀錄中心僅陳韋廷可以查看', 403)
    limit = int(request.args.get('limit') or 200)
    username = (request.args.get('username') or '').strip()
    entity_type = (request.args.get('entity_type') or '').strip()
    keyword = (request.args.get('q') or '').strip().lower()
    start_date = (request.args.get('start_date') or '').strip()
    end_date = (request.args.get('end_date') or '').strip()
    items = list_audit_trails(limit=max(limit * 4, 200))
    filtered = []
    for item in items:
        item = _decorate_audit_item(item)
        if username and username not in (item.get('username') or ''):
            continue
        if entity_type and entity_type not in (item.get('entity_type') or ''):
            continue
        created = (item.get('created_at') or '')[:10]
        if start_date and created and created < start_date:
            continue
        if end_date and created and created > end_date:
            continue
        hay = json.dumps(item, ensure_ascii=False).lower()
        if keyword and keyword not in hay:
            continue
        filtered.append(item)
        if len(filtered) >= limit:
            break
    return jsonify(success=True, items=filtered)

@app.route('/api/audit-trails/bulk-delete', methods=['POST'])
@login_required_json
def api_audit_trails_bulk_delete():
    if not _is_admin_user():
        return error_response('只有陳韋廷可以批量刪除操作紀錄', 403)
    data = request.get_json(silent=True) or {}
    ids = data.get('ids') or []
    count = _delete_rows_by_ids('audit_trails', ids)
    notify_sync_event(kind='refresh', module='today_changes', message='操作紀錄已批量刪除', extra={'count': count})
    return jsonify(success=True, deleted=count)

@app.route('/api/today-changes/bulk-delete', methods=['POST'])
@login_required_json
def api_today_changes_bulk_delete():
    if not _is_admin_user():
        return error_response('只有陳韋廷可以批量刪除今日異動', 403)
    data = request.get_json(silent=True) or {}
    ids = data.get('ids') or []
    count = _delete_rows_by_ids('logs', ids)
    notify_sync_event(kind='refresh', module='today_changes', message='今日異動已批量刪除', extra={'count': count})
    return jsonify(success=True, deleted=count, **_today_changes_payload())

@app.route('/api/customer-specs', methods=['GET'])
@login_required_json
def api_customer_specs():
    name = (request.args.get('name') or '').strip()
    return jsonify(success=True, items=get_customer_spec_stats(name, limit=int(request.args.get('limit') or 20)))

@app.route('/api/reports/export', methods=['GET'])
@login_required_json
def api_reports_export():
    report_type = (request.args.get('type') or 'inventory').strip()
    start_date = (request.args.get('start_date') or '').strip()
    end_date = (request.args.get('end_date') or '').strip()
    q = (request.args.get('q') or '').strip()

    if report_type == 'inventory':
        rows = inventory_summary()
        columns = [('客戶', 'customer_name'), ('商品', 'product_text'), ('總數量', 'qty'), ('已放倉庫', 'placed_qty'), ('未放倉庫', 'unplaced_qty'), ('位置', 'location'), ('操作人員', 'operator'), ('更新時間', 'updated_at')]
        name = '庫存總表.xlsx'
    elif report_type == 'orders':
        rows = get_orders()
        columns = [('客戶', 'customer_name'), ('商品', 'product_text'), ('數量', 'qty'), ('狀態', 'status'), ('操作人員', 'operator'), ('建立時間', 'created_at'), ('更新時間', 'updated_at')]
        name = '訂單總表.xlsx'
    elif report_type == 'shipping':
        rows = get_shipping_records(start_date or None, end_date or None, q)
        columns = [('客戶', 'customer_name'), ('商品', 'product_text'), ('數量', 'qty'), ('操作人員', 'operator'), ('出貨時間', 'shipped_at'), ('備註', 'note')]
        name = '出貨紀錄.xlsx'
    elif report_type == 'master_orders':
        rows = get_master_orders()
        columns = [('客戶', 'customer_name'), ('商品', 'product_text'), ('數量', 'qty'), ('操作人員', 'operator'), ('建立時間', 'created_at'), ('更新時間', 'updated_at')]
        name = '客戶總單.xlsx'
    elif report_type == 'unplaced':
        rows = [r for r in inventory_summary() if int(r.get('unplaced_qty') or 0) > 0]
        columns = [('客戶', 'customer_name'), ('商品', 'product_text'), ('未放倉庫數量', 'unplaced_qty'), ('總數量', 'qty'), ('位置', 'location'), ('更新時間', 'updated_at')]
        name = '未入倉商品.xlsx'
    elif report_type == 'warehouse':
        rows = []
        for cell in warehouse_get_cells():
            try:
                items = json.loads(cell.get('items_json') or '[]')
            except Exception:
                items = []
            if not items:
                rows.append({**cell, 'location': f"{cell.get('zone')}-{cell.get('column_index')}-{str(cell.get('slot_number')).zfill(2)}", 'customer_name': '', 'product_text': '', 'qty': 0})
            else:
                for it in items:
                    rows.append({**cell, 'location': f"{cell.get('zone')}-{cell.get('column_index')}-{str(cell.get('slot_number')).zfill(2)}", 'customer_name': it.get('customer_name') or '', 'product_text': it.get('product_text') or it.get('product') or '', 'qty': it.get('qty') or 0})
        columns = [('格位', 'location'), ('區域', 'zone'), ('欄', 'column_index'), ('格號', 'slot_number'), ('客戶', 'customer_name'), ('商品', 'product_text'), ('數量', 'qty'), ('備註', 'note'), ('更新時間', 'updated_at')]
        name = '倉庫位置表.xlsx'
    elif report_type == 'audit_trails':
        if current_username() != '陳韋廷':
            return error_response('操作紀錄僅陳韋廷可以匯出', 403)
        rows = []
        for item in list_audit_trails(limit=5000):
            decorated = _decorate_audit_item(item)
            rows.append(decorated)
        columns = [('時間', 'created_at'), ('操作者', 'username'), ('操作', 'action_label'), ('資料類型', 'entity_label'), ('資料鍵值', 'entity_key'), ('摘要', 'summary_text'), ('變更前', 'before_text'), ('變更後', 'after_text')]
        name = '操作紀錄.xlsx'
    elif report_type == 'customers':
        rows = get_customers(active_only=False)
        columns = [('客戶UID', 'customer_uid'), ('客戶名稱', 'name'), ('電話', 'phone'), ('地址', 'address'), ('區域', 'region'), ('特殊要求', 'notes'), ('封存', 'is_archived'), ('更新時間', 'updated_at')]
        name = '客戶資料.xlsx'
    else:
        return error_response('報表類型不存在')

    buf = export_rows_to_xlsx(report_type, rows, columns)
    return send_file(buf, as_attachment=True, download_name=name, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')



@app.route('/api/backups/download/<path:filename>', methods=['GET'])
@login_required_json
def api_backup_download(filename):
    safe_name = os.path.basename(filename)
    path = os.path.join('backups', safe_name)
    if not os.path.isfile(path):
        return error_response('找不到備份檔', 404)
    return send_file(path, as_attachment=True, download_name=safe_name)

@app.route('/api/backups/restore', methods=['POST'])
@login_required_json
def api_backup_restore():
    if current_username() != '陳韋廷':
        return error_response('權限不足', 403)
    data = request.get_json(silent=True) or {}
    filename = os.path.basename((data.get('filename') or '').strip())
    if not filename:
        return error_response('請選擇備份檔')
    path = os.path.join('backups', filename)
    if not os.path.isfile(path):
        return error_response('找不到備份檔', 404)
    if filename.endswith('.json'):
        payload = json.load(open(path, 'r', encoding='utf-8'))
        conn = get_db(); cur = conn.cursor()
        tables = ['users','inventory','orders','master_orders','shipping_records','corrections','image_hashes','logs','errors','warehouse_cells','customer_profiles','app_settings','customer_aliases','warehouse_recent_slots','audit_trails','todo_items']
        try:
            for table in tables:
                if table not in payload:
                    continue
                cur.execute(sql(f'DELETE FROM {table}'))
                rows = payload.get(table) or []
                for row in rows:
                    keys = list(row.keys())
                    if not keys:
                        continue
                    cols = ','.join(keys)
                    holders = ','.join(['?'] * len(keys))
                    cur.execute(sql(f'INSERT INTO {table}({cols}) VALUES ({holders})'), tuple(row[k] for k in keys))
            conn.commit()
        except Exception as e:
            conn.rollback()
            log_error('backup_restore', str(e))
            return error_response('還原失敗')
        finally:
            conn.close()
        notify_sync_event(kind='refresh', module='all', message='已還原備份')
        return jsonify(success=True)
    return error_response('目前只支援 JSON 備份還原')


@app.route('/api/undo-last', methods=['POST'])
@login_required_json
def api_undo_last():
    try:
        trails = list_audit_trails(limit=120)
        target = None
        for item in trails:
            if (item.get('username') or '') != current_username():
                continue
            if item.get('entity_type') == 'undo':
                continue
            if item.get('action_type') not in ('create','ship','move'):
                continue
            target = item
            break
        if not target:
            return error_response('目前沒有可還原的最近操作')
        conn = get_db()
        cur = conn.cursor()
        action_type = target.get('action_type')
        entity_type = target.get('entity_type')
        after_json = target.get('after_json') or {}
        entity_key = target.get('entity_key') or ''
        summary = ''
        if action_type == 'create' and entity_type == 'inventory':
            for it in after_json.get('items') or []:
                cur.execute(sql('SELECT id, qty FROM inventory WHERE customer_name = ? AND product_text = ? ORDER BY id DESC'), ((after_json.get('customer_name') or '').strip(), (it.get('product_text') or '').strip()))
                row = fetchone_dict(cur)
                if row:
                    cur.execute(sql('DELETE FROM inventory WHERE id = ?'), (row.get('id'),))
            summary = '已還原最近一次建立庫存'
        elif action_type == 'create' and entity_type == 'orders':
            for it in after_json.get('items') or []:
                cur.execute(sql('SELECT id FROM orders WHERE customer_name = ? AND product_text = ? ORDER BY id DESC'), ((after_json.get('customer_name') or '').strip(), (it.get('product_text') or '').strip()))
                row = fetchone_dict(cur)
                if row:
                    cur.execute(sql('DELETE FROM orders WHERE id = ?'), (row.get('id'),))
            summary = '已還原最近一次建立訂單'
        elif action_type == 'create' and entity_type == 'master_orders':
            for it in after_json.get('items') or []:
                cur.execute(sql('SELECT id FROM master_orders WHERE customer_name = ? AND product_text = ? ORDER BY id DESC'), ((after_json.get('customer_name') or '').strip(), (it.get('product_text') or '').strip()))
                row = fetchone_dict(cur)
                if row:
                    cur.execute(sql('DELETE FROM master_orders WHERE id = ?'), (row.get('id'),))
            summary = '已還原最近一次建立總單'
        elif action_type == 'ship':
            for item in after_json.get('breakdown') or []:
                for d in item.get('master_details') or []:
                    cur.execute(sql('UPDATE master_orders SET qty = qty + ?, updated_at = ? WHERE id = ?'), (int(d.get('qty') or 0), now(), int(d.get('id'))))
                for d in item.get('order_details') or []:
                    cur.execute(sql('UPDATE orders SET qty = qty + ?, updated_at = ? WHERE id = ?'), (int(d.get('qty') or 0), now(), int(d.get('id'))))
                for d in item.get('inventory_details') or []:
                    cur.execute(sql('UPDATE inventory SET qty = qty + ?, updated_at = ? WHERE id = ?'), (int(d.get('qty') or 0), now(), int(d.get('id'))))
                cur.execute(sql('SELECT id FROM shipping_records WHERE customer_name = ? AND product_text = ? ORDER BY id DESC'), ((after_json.get('customer_name') or '').strip(), (item.get('product_text') or '').strip()))
                ship_row = fetchone_dict(cur)
                if ship_row:
                    cur.execute(sql('DELETE FROM shipping_records WHERE id = ?'), (int(ship_row.get('id')),))
            summary = '已還原最近一次出貨'
        elif action_type == 'move' and entity_type == 'warehouse_cells':
            before_key = tuple((target.get('before_json') or {}).get('from_key') or [])
            to_key = tuple((target.get('after_json') or {}).get('to_key') or [])
            product_text = (target.get('after_json') or {}).get('product_text') or entity_key
            qty = int((target.get('after_json') or {}).get('qty') or 1)
            customer_name = (target.get('after_json') or {}).get('customer_name') or ''
            result = warehouse_move_item(to_key, before_key, product_text, qty, customer_name=customer_name)
            if not result.get('success'):
                conn.close()
                return error_response(result.get('error') or '還原倉庫移動失敗')
            summary = '已還原最近一次倉庫搬移'
        else:
            conn.close()
            return error_response('這筆操作暫不支援還原')
        conn.commit()
        conn.close()
        add_audit_trail(current_username(), 'undo', 'undo', entity_type, before_json=target, after_json={'message': summary})
        notify_sync_event(kind='refresh', module='all', message=summary)
        log_action(current_username(), summary)
        return jsonify(success=True, message=summary)
    except Exception as e:
        try:
            conn.rollback()
            conn.close()
        except Exception:
            pass
        log_error('undo_last', str(e))
        return error_response('還原上一筆失敗')


@app.route('/api/session/config', methods=['GET'])
@login_required_json
def api_session_config():
    return jsonify(success=True, idle_timeout_seconds=1800, startup_checks=STARTUP_CHECKS)

@app.route("/health")
@app.route("/api/health")
def health():
    return jsonify(success=True, status="ok", service="yuanxing", mode="native_device_only")

@app.route("/api/native-shell/config", methods=["GET"])
def api_native_shell_config():
    return jsonify(success=True, backend_url=request.host_url.rstrip("/"), allowed_origins=["capacitor://localhost", "http://localhost", "https://localhost", "ionic://localhost", "app://localhost", "null"], app_name="沅興木業")


# ==== FIX28：庫存 / 訂單 / 總單 / 出貨互通 API ====
def _fix28_table_for_source(source):
    s = (source or '').strip()
    mapping = {
        'inventory': ('inventory', '庫存'), '庫存': ('inventory', '庫存'),
        'orders': ('orders', '訂單'), 'order': ('orders', '訂單'), '訂單': ('orders', '訂單'),
        'master_order': ('master_orders', '總單'), 'master_orders': ('master_orders', '總單'), '總單': ('master_orders', '總單'),
    }
    return mapping.get(s)

def _fix28_get_row(table, item_id):
    conn = get_db(); cur = conn.cursor()
    cur.execute(sql(f"SELECT * FROM {table} WHERE id = ?"), (int(item_id),))
    row = fetchone_dict(cur)
    conn.close()
    return row

def _fix28_update_or_delete_source(table, item_id, move_qty):
    row = _fix28_get_row(table, item_id)
    if not row:
        return False, '找不到來源資料'
    current_qty = int(row.get('qty') or 0)
    move_qty = max(1, min(int(move_qty or current_qty), current_qty))
    conn = get_db(); cur = conn.cursor()
    if move_qty >= current_qty:
        cur.execute(sql(f"DELETE FROM {table} WHERE id = ?"), (int(item_id),))
    else:
        cur.execute(sql(f"UPDATE {table} SET qty = qty - ?, updated_at = ? WHERE id = ?"), (move_qty, now(), int(item_id)))
    conn.commit(); conn.close()
    return True, move_qty

def _fix28_update_item_api(table, item_id):
    try:
        row = _fix28_get_row(table, item_id)
        if not row:
            return error_response('找不到資料', 404)
        if request.method == 'GET':
            return jsonify(success=True, item=row)
        if request.method == 'DELETE':
            conn = get_db(); cur = conn.cursor()
            cur.execute(sql(f"DELETE FROM {table} WHERE id = ?"), (int(item_id),))
            conn.commit(); conn.close()
            add_audit_trail(current_username(), 'delete', table, str(item_id), before_json=row, after_json={})
            log_action(current_username(), f"刪除{('訂單' if table=='orders' else '總單')}商品 #{item_id}")
            notify_sync_event(kind='refresh', module=('orders' if table=='orders' else 'master_order'), message='商品已刪除', extra={'id': item_id})
            return jsonify(success=True, items=(get_orders() if table=='orders' else get_master_orders()))
        data = request.get_json(silent=True) or {}
        product_text = format_product_text_height2((data.get('product_text') or row.get('product_text') or '').strip())
        material = (data.get('material') if data.get('material') is not None else row.get('material') or '').strip().upper()
        product_code = clean_material_value(data.get('product_code') or material or '', product_text)
        customer_name = (data.get('customer_name') or row.get('customer_name') or '').strip()
        qty = int(data.get('qty') if data.get('qty') is not None else row.get('qty') or 0)
        if not product_text or not customer_name:
            return error_response('請輸入客戶與商品資料')
        if qty < 0:
            qty = 0
        conn = get_db(); cur = conn.cursor()
        cur.execute(sql(f"UPDATE {table} SET customer_name = ?, product_text = ?, product_code = ?, material = ?, qty = ?, operator = ?, updated_at = ? WHERE id = ?"), (customer_name, product_text, product_code, material, qty, current_username(), now(), int(item_id)))
        conn.commit(); conn.close()
        upsert_customer(customer_name, region=resolve_customer_region(customer_name, data.get('region')))
        add_audit_trail(current_username(), 'update', table, str(item_id), before_json=row, after_json={'customer_name': customer_name, 'product_text': product_text, 'material': material, 'qty': qty})
        log_action(current_username(), f"修改{('訂單' if table=='orders' else '總單')}商品 #{item_id}")
        notify_sync_event(kind='refresh', module=('orders' if table=='orders' else 'master_order'), message='商品已更新', extra={'id': item_id})
        return jsonify(success=True, items=(get_orders() if table=='orders' else get_master_orders()))
    except Exception as e:
        log_error('fix28_update_item_api', str(e))
        return error_response('商品操作失敗')

@app.route('/api/orders/<int:item_id>', methods=['GET', 'PUT', 'DELETE'])
@login_required_json
def api_fix28_order_item(item_id):
    return _fix28_update_item_api('orders', item_id)

@app.route('/api/master_orders/<int:item_id>', methods=['GET', 'PUT', 'DELETE'])
@login_required_json
def api_fix28_master_item(item_id):
    return _fix28_update_item_api('master_orders', item_id)

@app.route('/api/items/transfer', methods=['POST'])
@login_required_json
def api_fix28_items_transfer():
    try:
        data = request.get_json(silent=True) or {}
        source_info = _fix28_table_for_source(data.get('source'))
        if not source_info:
            return error_response('來源類型錯誤')
        source_table, source_label = source_info
        item_id = int(data.get('id') or 0)
        row = _fix28_get_row(source_table, item_id)
        if not row:
            return error_response('找不到來源商品', 404)
        current_qty = int(row.get('qty') or 0)
        qty = int(data.get('qty') or current_qty or 0)
        if qty <= 0:
            return error_response('數量必須大於 0')
        qty = min(qty, current_qty)
        target = (data.get('target') or '').strip()
        product_text = format_product_text_height2((row.get('product_text') or '').strip())
        material = (row.get('material') or ((row.get('product_code') or '') if (row.get('product_code') or '') != product_text else '')).strip()
        product_code = clean_material_value(row.get('product_code') or material or '', product_text)
        customer_name = (data.get('customer_name') or row.get('customer_name') or '').strip()
        item = {'product_text': product_text, 'product_code': product_code, 'material': material, 'qty': qty}
        target_label = ''
        result_payload = {}
        if target == 'inventory':
            save_inventory_item(product_text, product_code, qty, (data.get('location') or row.get('location') or '').strip(), customer_name, current_username(), f'from {source_table}', material)
            target_label = '庫存'
            ok, moved = _fix28_update_or_delete_source(source_table, item_id, qty)
            if not ok: return error_response(moved)
            result_payload['items'] = grouped_inventory()
        elif target == 'orders':
            if not customer_name: return error_response('請選擇客戶')
            upsert_customer(customer_name, region=resolve_customer_region(customer_name, data.get('region')))
            save_order(customer_name, [item], current_username(), (data.get('duplicate_mode') or 'merge').strip() or 'merge')
            target_label = '訂單'
            ok, moved = _fix28_update_or_delete_source(source_table, item_id, qty)
            if not ok: return error_response(moved)
            result_payload['items'] = get_orders()
        elif target in ('master_order', 'master_orders'):
            if not customer_name: return error_response('請選擇客戶')
            upsert_customer(customer_name, region=resolve_customer_region(customer_name, data.get('region')))
            save_master_order(customer_name, [item], current_username(), (data.get('duplicate_mode') or 'merge').strip() or 'merge')
            target_label = '總單'
            ok, moved = _fix28_update_or_delete_source(source_table, item_id, qty)
            if not ok: return error_response(moved)
            result_payload['items'] = get_master_orders()
        elif target == 'ship':
            if not customer_name: return error_response('請選擇客戶')
            upsert_customer(customer_name, region=resolve_customer_region(customer_name, data.get('region')))
            ship_result = ship_order(customer_name, [item], current_username(), allow_inventory_fallback=bool(data.get('allow_inventory_fallback')))
            if not ship_result.get('success'):
                return jsonify(ship_result), 400
            target_label = '出貨'
            result_payload.update(ship_result)
        else:
            return error_response('目標類型錯誤')
        add_audit_trail(current_username(), 'transfer', source_table, str(item_id), before_json={'source': source_label, **row}, after_json={'target': target_label, 'customer_name': customer_name, 'product_text': product_text, 'qty': qty})
        log_action(current_username(), f"{source_label}移到{target_label}：{customer_name} {product_text}x{qty}")
        notify_sync_event(kind='refresh', module='all', message=f'{source_label}已移到{target_label}', extra={'source': source_label, 'target': target_label, 'customer_name': customer_name, 'product_text': product_text, 'qty': qty})
        return jsonify(success=True, message=f'已從{source_label}移到{target_label}', customer_name=customer_name, target=target_label, **result_payload)
    except Exception as e:
        log_error('fix28_items_transfer', str(e))
        return error_response('互通操作失敗')

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 10000))
    app.run(host="0.0.0.0", port=port)
